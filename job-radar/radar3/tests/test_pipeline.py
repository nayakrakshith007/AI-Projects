"""Combine, workbook build (edits survive) and the runner with fake sources: ok, empty twice (broken), failing (keeps old jobs), lock."""
from datetime import date

import pytest
from openpyxl import load_workbook

from radar3 import build as build_mod
from radar3 import config, runner
from radar3.readers.common import make_job
from radar3.state import RunState

TODAY = date(2026, 10, 13)


def job(title, company, location, remote, source="A", url=None, posted=None):
    s = {"id": source, "name": f"Source {source}", "cls": "A"}
    return make_job(s, title, url or f"https://example.com/{source}/{title}/{company}".replace(" ", "-"), location=location, remote=remote,
                    company=company, posted=posted)


def test_combine_drops_not_open_to_india_dedupes_and_ranks():
    from radar3.combine import combine
    jobs = {
        "A": [job("Machine Learning Engineer", "Acme", "Worldwide", True, "A"),
              job("Machine Learning Engineer", "Acme", "USA Only", True, "A", url="https://example.com/us"),
              job("Data Scientist", "Beta", "Bengaluru", False, "A", posted="2026-10-10"),
              job("Data Engineer", "Gamma", "Remote", True, "A"),
              job("Data Analyst", "Delta", "Hyderabad", False, "A"),
              job("Office Manager", "Zeta", "Worldwide", True, "A")],
        "B": [job("Machine Learning Engineer", "Acme Inc", "Anywhere", True, "B")],           # same job on a second source
    }
    matched, unmatched, stats = combine(jobs)
    assert stats["not_open_to_india"] == 1 and stats["duplicates_removed"] == 1
    assert [j["title"] for j in matched] == ["Machine Learning Engineer", "Data Engineer", "Data Scientist", "Data Analyst"]   # tiers 0, 1, 2, 3
    assert [j["tier"] for j in matched] == [0, 1, 2, 3]
    assert len(matched[0]["sites"]) == 2 and matched[0]["source_ids"] == ["A", "B"]
    assert [j["title"] for j in unmatched] == ["Office Manager"]
    assert stats["remote_check_eligibility"] == 1


def test_sheet_names_are_valid_for_excel():
    for name in ["AI Architect / Solutions Architect", "GenAI / LLM Engineer", "Applied Scientist / Research Engineer"]:
        n = build_mod.sheet_name(name)
        assert len(n) <= 31 and "/" not in n


def test_build_workbook_and_keep_edits(tmp_path):
    path = tmp_path / "Job_Radar_V3.xlsx"
    jobs = {"A": [job("Machine Learning Engineer", "Acme", "Worldwide", True), job("Data Scientist", "Beta", "Remote", True)]}
    written, stats = build_mod.build(jobs, path=path, snapshot_path=tmp_path / "snap.json", today=TODAY)
    wb = load_workbook(written)
    assert wb.sheetnames[:2] == ["New today", "All"] and "Unmatched" in wb.sheetnames
    ws = wb["All"]
    heads = [c.value for c in ws[1]]
    assert heads[:16] == ["Roles", "Title", "Company", "Location", "Bengaluru", "Work mode", "Experience", "Salary", "Posted date", "Sites", "Skills",
                          "First seen", "Link", "Other links", "Status", "Notes"] and heads[16:] == ["Source type", "Check"]
    assert ws.cell(row=2, column=heads.index("Link") + 1).hyperlink.target.startswith("https://example.com/")
    assert ws.cell(row=3, column=heads.index("Check") + 1).value == "check eligibility"
    assert stats["new_today"] == 2
    wb.close()
    # the owner marks a job Applied and adds a note; the next day's build keeps them and the First seen date
    wb = load_workbook(path)
    ws = wb["All"]
    ws.cell(row=2, column=heads.index("Status") + 1).value = "Applied"
    ws.cell(row=2, column=heads.index("Notes") + 1).value = "called Tuesday"
    wb.save(path)
    wb.close()
    _, stats2 = build_mod.build(jobs, path=path, snapshot_path=None, today=date(2026, 10, 14))
    ws = load_workbook(path)["All"]
    assert (ws.cell(row=2, column=heads.index("Status") + 1).value, ws.cell(row=2, column=heads.index("Notes") + 1).value) == ("Applied", "called Tuesday")
    assert ws.cell(row=2, column=heads.index("First seen") + 1).value.date() == TODAY
    assert stats2["new_today"] == 0


# ---- runner with fake sources
class FakeReaders:
    def __init__(self):
        self.mode = {"S1": "ok", "S2": "empty", "S3": "boom"}

    def __call__(self, sid):
        def read(src, http):
            m = self.mode[sid]
            if m == "boom":
                raise RuntimeError("site down")
            return [] if m == "empty" else [job("Data Scientist", f"Co{sid}", "Worldwide", True, sid)]
        return read


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SOURCE_STATS_PATH", tmp_path / "stats.json")
    monkeypatch.setattr(config, "HISTORY_PATH", tmp_path / "history.jsonl")
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    fake = FakeReaders()
    monkeypatch.setattr(runner, "READERS", {f"r{i}": fake(f"S{i}") for i in (1, 2, 3)})
    sources = [{"id": f"S{i}", "name": f"Source S{i}", "company": "", "cls": "A", "reader": f"r{i}", "slug": "", "schedule": "daily", "enabled": "yes"} for i in (1, 2, 3)]
    return tmp_path, fake, sources


def _run(env, **kw):
    tmp, _, sources = env
    return runner.run("daily", sources=sources, http=object(), today=TODAY, state=RunState(tmp / "state.json"), out_dir=tmp / "sources",
                      excel_path=tmp / "wb.xlsx", **kw)


def test_runner_statuses_and_canary(env):
    tmp, fake, _ = env
    s1 = _run(env)
    st = RunState(tmp / "state.json")
    import json
    state = json.loads((tmp / "state.json").read_text())
    assert state["sources"]["S1"]["status"] == "done" and state["sources"]["S2"]["status"] == "empty" and state["sources"]["S3"]["status"] == "failed"
    assert s1["ok"] == 1 and set(s1["problems"]) == {"S2", "S3"}
    s2 = _run(env)                                                                     # empty a second time in a row: broken
    assert json.loads((tmp / "state.json").read_text())["sources"]["S2"]["status"] == "broken"
    assert (tmp / "wb.xlsx").exists() and s2["build"]["matched"] == 1


def test_failed_source_keeps_previous_jobs(env):
    tmp, fake, _ = env
    fake.mode["S3"] = "ok"
    _run(env)
    assert len(runner.load_source_jobs("S3", tmp / "sources")) == 1
    fake.mode["S3"] = "boom"
    summary = _run(env)
    assert len(runner.load_source_jobs("S3", tmp / "sources")) == 1                    # old jobs still there
    assert summary["build"]["matched"] == 2                                            # S1 and S3 (S2 never returned a job)


def test_lock_blocks_a_second_run(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    runner.acquire_lock()
    with pytest.raises(runner.AlreadyRunning):
        runner.acquire_lock()
    runner.release_lock()
    runner.acquire_lock()                                                              # free again
    runner.release_lock()


def test_stale_lock_from_dead_process_is_ignored(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    (tmp_path / "run.lock").write_text("999999")
    runner.acquire_lock()
    runner.release_lock()


def test_select_modes():
    from radar3.sources import select
    rows = [{"id": "D", "schedule": "daily", "enabled": "yes"}, {"id": "W", "schedule": "weekly", "enabled": "yes"}, {"id": "X", "schedule": "daily", "enabled": "no"}]
    assert [s["id"] for s in select(rows, "daily")] == ["D"]
    assert [s["id"] for s in select(rows, "auto", {}, TODAY)] == ["D", "W"]                      # a weekly source never read is due
    assert [s["id"] for s in select(rows, "auto", {"W": {"last_ok": "2026-10-10"}}, TODAY)] == ["D"]
    assert [s["id"] for s in select(rows, "auto", {"W": {"last_ok": "2026-10-05"}}, TODAY)] == ["D", "W"]
    assert [s["id"] for s in select(rows, "all")] == ["D", "W"]
    assert [s["id"] for s in select(rows, "daily", only=["W"])] == ["W"]


def test_saved_jobs_are_reclassified_when_loaded(tmp_path):
    """A rule change (e.g. US states) applies to saved reads without calling the sites again."""
    j = job("Data Scientist", "Acme", "Remote - California", True)
    j.update(tier=1, eligibility="unclear")                       # as an older rule set saved it
    runner.save_source_jobs("S9", [j], tmp_path)
    assert runner.load_source_jobs("S9", tmp_path)[0]["tier"] is None
    w = job("Data Scientist", "Acme", "2 Locations", False)
    w.update(tier=4, eligibility="unclear")                       # Workday's multi-location flag stays
    runner.save_source_jobs("S8", [w], tmp_path)
    assert runner.load_source_jobs("S8", tmp_path)[0]["tier"] == 4
