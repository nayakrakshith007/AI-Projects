"""The job-boards lane: V2's site workflows run from V3. The site runner and the site workbooks are faked; nothing opens a browser."""
import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from radar3 import boards, config, runner
from radar3.readers import READERS
from radar3.readers.common import make_job
from radar3.state import RunState

TODAY = date.today()
SRC = {"id": "J01", "name": "Naukri", "company": "Naukri", "cls": "B", "reader": "board_site", "slug": "naukri", "schedule": "daily", "enabled": "yes"}


def v2job(title="Senior Data Scientist", location="Bengaluru, Karnataka", url="https://www.naukri.com/job/1", posted=None, seen=None, remote=False, **kw):
    return {"title": title, "company": "Acme", "location": location, "more_locations": "", "experience": "5-8 Yrs", "salary": "", "work_mode": "", "skills": "python, sql",
            "posted_date": posted or TODAY, "first_seen": seen or TODAY, "role_hint": "Data Scientist", "job_type": "", "site": "naukri", "url": url,
            "bengaluru": "bengaluru" in location.lower(), "remote": remote, **kw}


def test_convert_makes_a_v3_job_with_location_priority_and_source_type():
    j = boards.convert(v2job(), SRC)
    assert j["source_type"] == "Job board" and j["source_id"] == "J01" and j["site"] == "Naukri" and j["job_id"] == j["url"]
    assert j["tier"] == 2 and j["bengaluru"] and j["eligibility"] == "india" and j["skills"] == "python, sql" and j["experience"] == "5-8 Yrs"
    assert boards.convert(v2job(location="Remote - India", remote=True), SRC)["tier"] == 0
    assert boards.convert(v2job(location="Pune", url="https://x/2"), SRC)["tier"] == 4


def test_convert_drops_jobs_older_than_the_age_limit():
    old = TODAY - timedelta(days=boards.MAX_AGE_DAYS + 1)
    assert boards.convert(v2job(posted=old, seen=old), SRC) is None
    assert boards.convert(v2job(posted=old, seen=TODAY), SRC) is not None           # newest of the two dates counts
    assert boards.convert(v2job(posted=None, seen=None), SRC) is not None


@pytest.mark.parametrize("lane,v3", [("ok", "done"), ("finished with errors", "done"), ("needs you", "needs you"), ("stopped", "stopped"), ("stalled", "failed"),
                                     ("timed out", "failed"), ("failed", "failed"), ("skipped", "failed")])
def test_status_mapping(lane, v3):
    assert boards.status_for({"status": lane, "problems": []})[0] == v3


def test_status_note_names_the_problem_and_the_second_try():
    status, note = boards.status_for({"status": "needs you", "problems": ["!! human check on page 3"], "retried": True, "first_status": "stalled"})
    assert "needs you" in note and "human check" in note and "second try after stalled" in note


class FakeSites:
    """Stands in for pipeline.runner.run_sites: reports each site to the callbacks like the real one."""
    def __init__(self, outcome):
        self.outcome, self.calls = outcome, []

    def __call__(self, sites, lanes=None, on_done=None, logs_dir=None, on_start=None, control=None, on_note=None, **kw):
        self.calls.append([s.name for s in sites])
        for s in sites:
            on_start(s.name, logs_dir / "x.log", 0)
            on_done(dict({"site": s.name, "status": "ok", "minutes": 1.5, "new": 4, "total": 40, "problems": []}, **self.outcome.get(s.name, {})))
        return []


def test_run_boards_reports_each_site_live_and_converts_jobs(tmp_path):
    state = RunState(tmp_path / "state.json")
    state.begin([SRC], "all")
    res = boards.run_boards([SRC], state, run_sites=FakeSites({}), read_site=lambda site: [v2job(), v2job(url="https://www.naukri.com/job/2", title="ML Engineer")], logs_dir=tmp_path)
    assert res["J01"]["status"] == "done" and len(res["J01"]["jobs"]) == 2 and res["J01"]["read"] is True
    assert state.data["sources"]["J01"]["status"] == "done" and state.data["sources"]["J01"]["jobs"] == 40         # the live row shows the site workbook total


def test_run_boards_survives_an_unreadable_workbook(tmp_path):
    state = RunState(tmp_path / "state.json")
    state.begin([SRC], "all")

    def broken(site):
        raise OSError("workbook locked")
    res = boards.run_boards([SRC], state, run_sites=FakeSites({}), read_site=broken, logs_dir=tmp_path)
    assert res["J01"]["read"] is False and res["J01"]["jobs"] == [] and "not readable" in res["J01"]["note"]


# ---- the whole run: fast API source + a job board side by side
@pytest.fixture
def env(tmp_path, monkeypatch):
    for name in ("SOURCE_STATS_PATH", "HISTORY_PATH", "HEALTH_PATH", "RUNS_PATH", "LAST_RUN_PATH", "SNAPSHOT_PATH"):
        monkeypatch.setattr(config, name, tmp_path / (name.lower() + ".json"))
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(config, "EXCEL_PATH", tmp_path / "wb.xlsx")
    monkeypatch.setitem(READERS, "fastr", lambda src, http: [make_job({"id": "A1", "name": "Feed", "cls": "A"}, "ML Engineer", "https://feed.test/1", location="Worldwide", remote=True, company="Beta")])
    api = {"id": "A1", "name": "Feed", "company": "", "cls": "A", "reader": "fastr", "slug": "", "schedule": "daily", "enabled": "yes"}
    return tmp_path, [api, SRC]


def go(tmp, sources, outcome, jobs, boards_on=True):
    return runner.run("daily", sources=sources, http=object(), today=TODAY, state=RunState(tmp / "state.json"), out_dir=tmp / "sources", excel_path=config.EXCEL_PATH,
                      use_lock=False, boards=boards_on, board_runner=FakeSites(outcome), board_reader=lambda site: jobs)


def test_run_reads_both_kinds_and_the_workbook_has_both(env):
    tmp, sources = env
    summary = go(tmp, sources, {}, [v2job(), v2job(url="https://www.naukri.com/job/2", title="Data Engineer")])
    assert summary["sources_run"] == 2 and summary["ok"] == 2 and summary["problems"] == {}
    ws = load_workbook(tmp / "wb.xlsx")["All"]
    heads = [c.value for c in ws[1]]
    kinds = {ws.cell(row=r, column=heads.index("Source type") + 1).value for r in range(2, ws.max_row + 1)}
    assert kinds == {"Feed", "Job board"} and ws.max_row == 4
    assert summary["build"]["by_source"] == {"A1": 1, "J01": 2}


def test_boards_can_be_left_out(env):
    tmp, sources = env
    summary = go(tmp, sources, {}, [v2job()], boards_on=False)
    assert summary["sources_run"] == 1 and "J01" not in json.loads((tmp / "state.json").read_text())["sources"]


def test_a_human_check_keeps_what_the_site_read_and_is_reported(env):
    tmp, sources = env
    summary = go(tmp, sources, {"naukri": {"status": "needs you", "problems": ["!! human check: pass it by hand"]}}, [v2job()])
    assert summary["problems"]["J01"]["status"] == "needs you" and "human check" in summary["problems"]["J01"]["note"]
    assert summary["build"]["by_source"].get("J01") == 1                                  # the partial read is still in the workbook


def test_a_failed_board_keeps_its_previous_list(env):
    tmp, sources = env
    go(tmp, sources, {}, [v2job(), v2job(url="https://www.naukri.com/job/2", title="Data Engineer")])
    summary = go(tmp, sources, {"naukri": {"status": "stalled", "problems": ["no output"]}}, [])
    assert summary["problems"]["J01"]["status"] == "failed"
    assert len(runner.load_source_jobs("J01", tmp / "sources")) == 2 and summary["build"]["by_source"].get("J01") == 2


def test_a_crashing_boards_lane_never_takes_the_run_down(env):
    tmp, sources = env

    def boom(*a, **k):
        raise RuntimeError("chrome missing")
    summary = runner.run("daily", sources=sources, http=object(), today=TODAY, state=RunState(tmp / "state.json"), out_dir=tmp / "sources", excel_path=config.EXCEL_PATH,
                         use_lock=False, board_runner=boom, board_reader=lambda s: [])
    assert summary["ok"] == 1 and summary["problems"]["J01"]["status"] == "failed" and "RuntimeError" in summary["problems"]["J01"]["note"]
