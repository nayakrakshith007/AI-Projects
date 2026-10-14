import json
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from common.inputs import Query
from hirist import config, excel
from hirist import search as S
from hirist.parse import (pages_for, parse_search_json, parse_tag_page, posted_to_date, split_title, title_matches,
                          total_from_title)

FIX = Path(__file__).parent / "fixtures"
TAG_HTML = (FIX / "tag_page.html").read_text(encoding="utf-8")
SEARCH_JSON = json.loads((FIX / "search.json").read_text(encoding="utf-8"))
TODAY = date(2026, 10, 11)


# ---------------------------------------------------------------- parsing
def test_tag_page_cards():
    jobs = parse_tag_page(TAG_HTML, TODAY)
    assert len(jobs) == 3
    first = jobs[0]
    assert first["id"] == "1677764" and first["url"] == "https://www.hirist.tech/j/indusind-bank-vice-president-data-science-1677764"
    assert (first["company"], first["title"]) == ("IndusInd Bank", "Vice President - Data Science")
    assert (first["experience"], first["experience_min"], first["experience_max"]) == ("10-12 years", 10, 12)
    assert first["salary"] == "50-55 LPA" and first["location"] == "Mumbai" and "Machine Learning" in first["skills"]
    assert first["posted_text"] == "1 day ago" and first["posted_date"] == date(2026, 10, 10) and first["source"] == "tag"
    assert jobs[2]["location"] == "Gurgaon, Gurugram" and jobs[2]["posted_date"] == date(2026, 10, 4)      # "1 week ago"


def test_total_from_title_and_pages():
    assert total_from_title(TAG_HTML) == 887 and pages_for(887) == 45 and pages_for(20) == 1 and pages_for(None) is None


def test_search_json_has_exact_fields_and_no_recruiter_details():
    jobs = parse_search_json(SEARCH_JSON)
    assert len(jobs) == 2
    a, b = jobs
    assert a["url"].endswith("-1676619") and a["id"] == "1676619" and a["source"] == "search"
    assert (a["company"], a["title"], a["designation"]) == ("Pattem Digital Technologies", "Forward Deployed Engineer - Full Stack Development", "Forward Deployed Engineer")
    assert (a["experience"], a["experience_min"], a["experience_max"]) == ("5-7 years", 5, 7)
    assert a["posted_at"] == "2026-10-05T13:28+05:30" and a["posted_date"] == date(2026, 10, 5)
    assert a["applicants"] == 228 and a["company_rating"] == 3.3 and a["salary"] == ""            # salary was hidden
    assert b["salary"] == "20-28 LPA" and b["location"] == "Bangalore, Chennai, Hyderabad" and b["bengaluru"] == "Yes"
    assert all("recruiter" not in j and "recruiter_name" not in j for j in jobs)


def test_title_split_and_posted_label():
    assert split_title("Acme - Data Scientist - NLP", "Acme Corp") == ("Acme Corp", "Data Scientist - NLP")
    assert split_title("Forward Deployment Engineer - Generative AI", "HiringBlaze") == ("HiringBlaze", "Forward Deployment Engineer - Generative AI")
    assert posted_to_date("Posted 3 days ago", TODAY) == date(2026, 10, 8) and posted_to_date("Posted today", TODAY) == TODAY


def test_title_matches_is_a_rough_relevance_check():
    assert title_matches({"title": "Senior Data Scientist", "designation": "", "skills": []}, "Data Scientist")
    assert title_matches({"title": "Lead", "designation": "", "skills": ["Machine Learning Engineer"]}, "Machine Learning Engineer")
    assert not title_matches({"title": "Sales Executive", "designation": "", "skills": ["CRM"]}, "Data Scientist")


# ---------------------------------------------------------------- urls and window
def test_posting_days_and_urls():
    assert [config.posting_days(h) for h in (12, 24, 33, 48, 81, 100, 168, 400, 0)] == [1, 1, 3, 3, 7, 7, 7, 7, 7]    # capped at a week
    assert config.tag_url("data-scientist", 7, 1) == "https://www.hirist.tech/k/data-scientist-jobs?posting=7"
    assert config.tag_url("data-scientist", 7, 3).endswith("posting=7&page=3")
    first, second = config.search_url("ai-engineer", 7, 1), config.search_url("ai-engineer", 7, 2)
    assert "/search/ai-engineer-jobs?" in first and "posting=7" in first and "page=0" in first and "page=2" in second


def test_in_window_exact_and_by_label():
    now = datetime(2026, 10, 12, 9, 0)
    assert S.in_window({"posted_at": "2026-10-11T10:00+05:30", "posted_date": date(2026, 10, 11)}, 33, now)
    assert not S.in_window({"posted_at": "2026-10-10T20:00+05:30", "posted_date": date(2026, 10, 10)}, 33, now)     # 37 h ago
    assert S.in_window({"posted_at": None, "posted_date": date(2026, 10, 11)}, 33, now)
    assert not S.in_window({"posted_at": None, "posted_date": date(2026, 10, 10)}, 33, now)
    assert S.in_window({"posted_at": None, "posted_date": None}, 33, now)


# ---------------------------------------------------------------- planning
def test_plan_builds_search_and_tag_lists():
    tags = S.load_tags()
    queries = [Query("AI Engineer", "AI Engineer", "India", 5), Query("AI Engineer", "LLM Engineer", "India", 5),
               Query("Forward Deployed Engineer", "Forward Deployed Engineer", "India", 5), Query("DS", "Data Scientist", "India", 5)]
    targets, notes = S.plan(queries, tags)
    pairs = [(t.method, t.slug) for t in targets]
    assert ("search", "ai-engineer") in pairs and ("search", "forward-deployed-engineer") in pairs         # search covers every keyword
    assert [p for p in pairs if p[0] == "tag"].count(("tag", "llm")) == 1                                    # shared tag is read once
    assert ("tag", "artificial-intelligence") in pairs and ("tag", "data-scientist") in pairs
    assert not any(t.method == "tag" and t.slug in ("-", "forward-deployed-engineer") for t in targets)
    assert any("Forward Deployed Engineer" in n and "no Hirist tag page" in n for n in notes)
    only_search, _ = S.plan(queries, tags, methods=("search",))
    assert {t.method for t in only_search} == {"search"}


# ---------------------------------------------------------------- the shared page queue (fake Hirist, no Chrome)
def mk(jid, title="Senior Data Scientist", exp=(5, 8), hours_ago=10, source="search", skills=("Python",)):
    posted = datetime.now() - timedelta(hours=hours_ago)
    lo, hi = exp
    return {"id": str(jid), "url": f"https://www.hirist.tech/j/job-{jid}", "title": title, "company": "Acme", "designation": "",
            "experience": f"{lo}-{hi} years" if hi else "", "experience_min": lo, "experience_max": hi, "locations": ["Bangalore"],
            "location": "Bangalore", "bengaluru": "Yes", "work_mode": "", "salary": "", "skills": list(skills), "posted_text": "",
            "posted_date": posted.date(), "posted_at": posted.isoformat(timespec="minutes") if source == "search" else None,
            "applicants": 5, "views": 50, "company_rating": 3.5, "company_reviews": 10, "source": source}


class _FakeSession:
    opened = 0

    def __init__(self, **kw):
        _FakeSession.opened += 1

    def page(self):
        return self

    def close(self):
        pass


def _patch(monkeypatch, lists):
    """lists: {(method, slug): [pages...]}, each page is a list of jobs. Returns the list of (method, slug, page) requested."""
    _FakeSession.opened = 0
    requested, lock = [], threading.Lock()

    def read_page(pg, t, days, page):
        with lock:
            requested.append((t.method, t.slug, page))
        time.sleep(0.01)
        pages = lists.get((t.method, t.slug), [])
        if not pages or page > len(pages):
            return [], (len(pages) or None)
        return pages[page - 1], len(pages)

    monkeypatch.setattr(S, "read_page", read_page)
    monkeypatch.setattr(S, "Session", _FakeSession)
    monkeypatch.setattr(S, "_pause", lambda rng: None)
    return requested


TAGS = {"data scientist": ["data-scientist"], "forward deployed engineer": ["-"]}


def test_pages_are_shared_and_each_read_once(monkeypatch):
    lists = {("search", "data-scientist"): [[mk(10 * p + i) for i in range(3)] for p in range(1, 6)]}
    requested = _patch(monkeypatch, lists)
    jobs, err, stats = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=3, methods=("search",))
    assert err is None and sorted(p for m, s, p in requested) == [1, 2, 3, 4, 5] and len(jobs) == 15
    assert stats["pages"] == 5 and _FakeSession.opened >= 1 and stats["posting_days"] == 7


def test_both_methods_are_merged_and_compared(monkeypatch):
    shared, only_s, only_t = mk(1), mk(2), mk(3, source="tag")
    lists = {("search", "data-scientist"): [[shared, only_s]],
             ("tag", "data-scientist"): [[dict(shared, source="tag", posted_at=None), only_t, mk(4, title="Sales Manager", source="tag", skills=("CRM",))]]}
    _patch(monkeypatch, lists)
    jobs, err, stats = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=2)
    by = {j["id"]: j for j in jobs}
    assert err is None and set(by) == {"1", "2", "3", "4"}
    assert by["1"]["methods"] == ["search", "tag"] and by["1"]["posted_at"]                     # found by both; keeps the exact time
    assert by["2"]["methods"] == ["search"] and by["3"]["methods"] == ["tag"]
    assert by["4"]["title_match"] == ""                                                         # the sales job does not look like the role
    row = stats["compare"][0]
    assert (row["search"], row["tag"], row["both"], row["only_search"], row["only_tag"], row["union"]) == (2, 3, 1, 1, 2, 4)
    assert row["tag_match_title"] == 2 and row["union_match_title"] == 3


def test_window_and_experience_rules(monkeypatch):
    lists = {("search", "data-scientist"): [[mk(1, hours_ago=5), mk(2, hours_ago=400), mk(3, exp=(0, 2)), mk(4, exp=(0, 0)), mk(5, exp=(3, 6))]]}
    _patch(monkeypatch, lists)
    jobs, _, stats = S.search([Query("DS", "Data Scientist", "India", 5)], 168, tags=TAGS, workers=1, methods=("search",))
    assert {j["id"] for j in jobs} == {"1", "4", "5"}          # 2 is too old, 3 asks for 0-2 years; no stated experience (4) is kept; 3-6 reaches 5
    assert stats["outside_window"] == 1 and stats["dropped_experience"] == 1


def test_human_check_stops_every_chrome(monkeypatch):
    from hirist.session import HumanCheckPending
    requested = _patch(monkeypatch, {("search", "data-scientist"): [[mk(p * 10 + i) for i in range(3)] for p in range(1, 30)]})
    real = S.read_page

    def read_page(pg, t, days, page):
        if page == 4:
            raise HumanCheckPending("Hirist is asking for a human check")
        return real(pg, t, days, page)

    monkeypatch.setattr(S, "read_page", read_page)
    jobs, err, _ = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=3, methods=("search",))
    assert "human check" in err and "check_hirist.bat" in err and len(requested) < 25 and jobs


def test_all_chromes_failing_does_not_hang(monkeypatch):
    _patch(monkeypatch, {})

    def boom(pg, t, days, page):
        raise RuntimeError("Chrome crashed")

    monkeypatch.setattr(S, "read_page", boom)
    jobs, err, _ = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=2, methods=("search",))
    assert jobs == [] and "RuntimeError" in err


# ---------------------------------------------------------------- Excel
def test_excel_adds_dedupes_keeps_user_fields_and_fills_details(tmp_path):
    path = tmp_path / "Hirist.xlsx"
    tag_job = dict(mk(7, source="tag"), role="Data Scientist", term="Data Scientist", methods=["tag"], title_match="Yes")
    assert excel.save([tag_job], path, TODAY)[1:] == (1, 0, 1)
    wb = load_workbook(path)
    wb["Hirist"].cell(row=2, column=excel.col("Notes")).value = "call them"
    wb.save(path)
    richer = dict(tag_job, methods=["search", "tag"], applicants=228, posted_at="2026-10-05T13:28+05:30")
    _, n_new, n_upd, total = excel.save([richer], path, date(2026, 10, 12))
    assert (n_new, n_upd, total) == (0, 1, 1)
    ws = load_workbook(path)["Hirist"]
    get = lambda name: ws.cell(row=2, column=excel.col(name)).value
    assert get("Applicants") == 228 and get("Found by") == "search + tag" and get("Notes") == "call them" and get("Role") == "Data Scientist"
    assert ws.cell(row=2, column=excel.col("Link")).hyperlink.target == tag_job["url"]
