from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from cutshort import config, excel
from cutshort import search as S
from cutshort.parse import experience, html_to_text, parse_page, posted_ist
from cutshort.session import HumanCheckPending

FIX = Path(__file__).parent / "fixtures"
PAGE = (FIX / "category_page.html").read_text(encoding="utf-8")
NOW = datetime(2026, 10, 11, 12, 0)


# ---------------------------------------------------------------- parsing
def test_page_jobs_and_total():
    jobs, total = parse_page(PAGE)
    assert total == 330 and len(jobs) == 5
    first = jobs[0]
    assert first["title"] == "Data Scientist – Agentic AI" and first["company"] == "VY SYSTEMS PRIVATE LIMITED"
    assert first["url"].endswith("VY-SYSTEMS-PRIVATE-LIMITED-Z13Xrplb") and first["id"] == "6ac5c90f6599b199749980a0"
    assert (first["experience"], first["experience_min"], first["experience_max"]) == ("5-8 years", 5, 8)
    assert first["salary"] == "₹2L - ₹15L / yr" and first["location"] == "Bengaluru (Bangalore), Hyderabad" and first["bengaluru"] == "Yes"
    assert first["work_mode"] == "Onsite" and first["job_type"] == "Full time" and first["posted_text"] == "7 Oct 2026"
    assert first["skills"][:3] == ["Data Science", "Python", "Machine Learning (ML)"]
    assert first["company_type"] == "Services" and first["company_size"] == "100-1000" and first["company_stage"] == "Profitable"
    assert first["company_website"] == "https://vysystems.com" and first["company_founded"] == 2000
    assert "5+ years in Data Science" in first["description"] and "<" not in first["description"]


def test_posted_time_is_indian_time():
    assert posted_ist("2026-10-07T04:22:39.818Z") == datetime(2026, 10, 7, 9, 52, 39, 818000)
    assert posted_ist("2026-10-07T20:00:00Z").date() == date(2026, 10, 8)          # after 18:30 UTC it is already the next day in India
    assert posted_ist(None) is None and posted_ist("nonsense") is None
    jobs, _ = parse_page(PAGE)
    assert jobs[0]["posted_date"] == date(2026, 10, 7) and jobs[0]["posted_at"] == "2026-10-07T09:52"


def test_remote_and_internship_and_staffing_jobs():
    jobs = {j["id"]: j for j in parse_page(PAGE)[0]}
    intern = jobs["6ab675c7bff5b7d1661d4e6f"]
    assert intern["work_mode"] == "Remote" and intern["job_types"] == ["internship"] and intern["experience"] == "" and intern["bengaluru"] == ""
    friendly = jobs["6aa28b01369e1cd8f328d7ac"]
    assert friendly["work_mode"] == "Onsite or remote" and friendly["location"] == "Remote, Noida"
    staffing = jobs["6ab366843339e4a13666f026"]
    assert staffing["hiring_for_client"] == "Yes" and staffing["company_website"] == "" and staffing["company_founded"] is None


def test_no_recruiter_name_is_kept():
    for j in parse_page(PAGE)[0]:
        assert "createdBy" not in j and "posted_by" not in j


def test_page_without_jobs_gives_nothing():
    assert parse_page("<html><body>Just a moment...</body></html>") == ([], None)
    assert parse_page('<script id="__NEXT_DATA__" type="application/json">{"props":{}}</script>') == ([], None)


def test_experience():
    assert experience({"min": 5, "max": 8}) == (5, 8, "5-8 years")
    assert experience({"min": 8, "max": 8}) == (8, 8, "8+ years")
    assert experience({"min": 0, "max": 0}) == (None, None, "") and experience(None) == (None, None, "")
    assert experience({"min": 3, "max": 0}) == (3, 3, "3+ years")


def test_html_to_text_keeps_structure():
    text = html_to_text("<p>Role</p><ul><li>One</li><li>Two &amp; three</li></ul><p>End<br/>Line</p>")
    assert text.splitlines()[0] == "Role" and "- One" in text and "- Two & three" in text and "End\nLine" in text
    assert len(html_to_text("<p>" + "x" * 50000 + "</p>")) == config.DESCRIPTION_LIMIT


# ---------------------------------------------------------------- window and rules
def test_window_uses_exact_time():
    job = {"posted_at": "2026-10-10T09:00"}
    assert S.in_window(job, 48, NOW) and not S.in_window(job, 24, NOW)
    assert S.in_window(job, 0, NOW) and S.in_window({"posted_at": None}, 24, NOW)


def test_keep_rules_count_why_jobs_drop():
    jobs = {j["id"]: j for j in parse_page(PAGE)[0]}
    counts = {"outside_window": 0, "dropped_internship": 0, "dropped_experience": 0}
    kept = [i for i, j in jobs.items() if S.keep(j, 168, 5, counts, NOW)]
    # 5-8 yrs and 3-10 yrs (top reaches 5) are inside 7 days; the other three were posted 10, 23 and 25 Sep
    assert kept == ["6ac5c90f6599b199749980a0", "6ac34a733f05bb8045d8b85d"]
    assert counts == {"outside_window": 3, "dropped_internship": 0, "dropped_experience": 0}      # 23 Sep, the internship (25 Sep) and 10 Sep
    wide = {"outside_window": 0, "dropped_internship": 0, "dropped_experience": 0}
    kept = [i for i, j in jobs.items() if S.keep(j, 800, 5, wide, NOW)]
    assert "6ab675c7bff5b7d1661d4e6f" not in kept and wide["dropped_internship"] == 1
    assert "6aa28b01369e1cd8f328d7ac" in kept                       # 2-5 years: the top of the range reaches 5
    assert wide["dropped_experience"] == 0
    strict = {"outside_window": 0, "dropped_internship": 0, "dropped_experience": 0}
    assert "6aa28b01369e1cd8f328d7ac" not in [i for i, j in jobs.items() if S.keep(j, 800, 6, strict, NOW)] and strict["dropped_experience"] >= 1


# ---------------------------------------------------------------- the run (Chrome faked)
class FakeSession:
    pages: dict = {}
    closed = 0

    def page(self):
        return self

    def close(self):
        FakeSession.closed += 1


def _run(monkeypatch, outcomes, hours=720):
    """outcomes: slug -> (jobs, total) or an exception to raise."""
    jobs_by_slug = parse_page(PAGE)[0]

    def fake_read(pg, lst):
        out = outcomes[lst.slug]
        if isinstance(out, Exception):
            raise out
        return [dict(j) for j in out[0]], out[1]

    monkeypatch.setattr(S, "Session", FakeSession)
    monkeypatch.setattr(S, "read_list", fake_read)
    monkeypatch.setattr(S, "_pause", lambda rng: None)
    lists = [config.JobList("a", "List A"), config.JobList("b", "List B"), config.JobList("c", "List C")]
    return jobs_by_slug, lists


def test_job_on_two_lists_is_kept_once_with_both_names(monkeypatch):
    jobs, lists = _run(monkeypatch, {"a": None, "b": None})
    outcomes = {"a": (jobs[:2], 100), "b": (jobs[:1], 40)}
    jobs, lists = _run(monkeypatch, outcomes)
    found, error, stats = S.search(720, 5, lists[:2])
    assert error is None and len(found) == 2 and stats["pages"] == 2 and stats["read"] == 3
    both = next(j for j in found if j["id"] == "6ac5c90f6599b199749980a0")
    assert both["lists"] == ["List A", "List B"]
    assert [r["kept"] for r in stats["reports"]] == [2, 1]


def test_human_check_stops_the_run_and_keeps_what_was_read(monkeypatch):
    jobs, lists = _run(monkeypatch, {})
    outcomes = {"a": (jobs[:2], 50), "b": HumanCheckPending("Cutshort is asking for a human check"), "c": (jobs[:1], 10)}
    jobs, lists = _run(monkeypatch, outcomes)
    found, error, stats = S.search(720, 5, lists)
    assert "human check" in error and len(found) == 2 and stats["pages"] == 1            # list C is never opened


def test_an_error_on_one_list_does_not_stop_the_others(monkeypatch):
    jobs, lists = _run(monkeypatch, {})
    outcomes = {"a": RuntimeError("boom"), "b": (jobs[:2], 50), "c": (jobs[:1], 10)}
    jobs, lists = _run(monkeypatch, outcomes)
    found, error, stats = S.search(720, 5, lists)
    assert "RuntimeError" in error and len(found) == 2 and stats["pages"] == 2


def test_empty_list_is_reported(monkeypatch):
    jobs, lists = _run(monkeypatch, {})
    jobs, lists = _run(monkeypatch, {"a": ([], None), "b": (jobs[:1], 10)})
    found, error, stats = S.search(720, 5, lists[:2])
    assert "a: no jobs" in error and len(found) == 1


# ---------------------------------------------------------------- excel
def test_excel_merge_keeps_user_columns_and_grows_lists(tmp_path):
    path = tmp_path / "Cutshort.xlsx"
    jobs = parse_page(PAGE)[0][:2]
    for j in jobs:
        j["lists"] = ["List A"]
    p, n_new, n_upd, total = excel.save(jobs, path, today=date(2026, 10, 9))
    assert (n_new, n_upd, total) == (2, 0, 2)
    wb = load_workbook(path)
    ws = wb[excel.SHEET]
    assert [c.value for c in ws[1]] == excel.HEADERS
    row = excel.HEADERS.index("Status") + 1
    ws.cell(row=2, column=row).value = "Applied"
    ws.cell(row=2, column=excel.HEADERS.index("Notes") + 1).value = "call Monday"
    wb.save(path)
    again = parse_page(PAGE)[0][:3]
    for j in again:
        j["lists"] = ["List B"]
    again[0]["salary"] = "₹3L - ₹16L / yr"
    _, n_new, n_upd, total = excel.save(again, path, today=date(2026, 10, 11))
    assert (n_new, n_upd, total) == (1, 2, 3)
    ws = load_workbook(path)[excel.SHEET]
    rows = {r[excel.HEADERS.index("Title")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    kept = rows["Data Scientist – Agentic AI"]
    assert kept[excel.HEADERS.index("Status")] == "Applied" and kept[excel.HEADERS.index("Notes")] == "call Monday"
    assert kept[excel.HEADERS.index("Salary")] == "₹3L - ₹16L / yr"                      # refreshed
    assert kept[excel.HEADERS.index("Found in")] == "List A, List B"                      # only grows
    assert kept[excel.HEADERS.index("First seen")].date() == date(2026, 10, 9)
    link = ws.cell(row=2, column=excel.HEADERS.index("Link") + 1)
    assert link.value == "Open" and link.hyperlink.target.startswith("https://cutshort.io/job/")


def test_excel_open_in_excel_writes_a_copy(tmp_path, monkeypatch):
    path = tmp_path / "Cutshort.xlsx"
    real = excel._write

    def locked(rows, p):
        if p == path:
            raise PermissionError
        real(rows, p)

    monkeypatch.setattr(excel, "_write", locked)
    jobs = parse_page(PAGE)[0][:1]
    jobs[0]["lists"] = ["List A"]
    written, *_ = excel.save(jobs, path)
    assert written != path and written.name.startswith("Cutshort_") and written.exists()
