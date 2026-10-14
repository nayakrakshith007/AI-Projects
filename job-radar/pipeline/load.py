"""Read every site workbook into one common job format. Header-driven, so a site may add or move columns."""
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from common.geo import is_bengaluru

from . import config

# common field -> the column names the sites use for it (first one found wins)
FIELDS = {
    "title": ("Title",), "company": ("Company",), "location": ("Location",), "more_locations": ("More locations",),
    "experience": ("Experience",), "salary": ("Salary",), "work_mode": ("Work mode",), "skills": ("Skills",),
    "posted_date": ("Posted date",), "first_seen": ("First seen",), "role_hint": ("Role",), "job_type": ("Job type",),
}


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def _link(ws, row_no: int, col: int) -> str:
    cell = ws.cell(row=row_no, column=col)
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    return str(cell.value) if cell.value and str(cell.value).startswith("http") else ""


def read_site(site: config.Site, path: Path | None = None) -> list[dict]:
    """All jobs in one site's workbook, in the common format. A missing workbook gives an empty list."""
    path = path or site.workbook
    if not path.exists():
        return []
    wb = load_workbook(path)                       # not read-only: hyperlinks are needed
    try:
        if site.sheet not in wb.sheetnames:
            return []
        ws = wb[site.sheet]
        names = [c.value for c in ws[1]]
        index = {n: i for i, n in enumerate(names)}
        link_col = index.get("Link")
        if link_col is None:
            return []
        jobs = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            url = _link(ws, row_no, link_col + 1)
            if not url:
                continue
            job = {"site": site.name, "url": url}
            for field, cols in FIELDS.items():
                v = next((row[index[c]] for c in cols if c in index and row[index[c]] not in (None, "")), "")
                job[field] = _date(v) if field in ("posted_date", "first_seen") else (str(v).strip() if v != "" else "")
            job["title"] = job["title"] or ""
            if not job["title"]:
                continue
            place = f'{job["location"]} {job["more_locations"]}'
            job["bengaluru"] = bool(is_bengaluru(place))
            mode = f'{job["work_mode"]} {job["location"]}'.lower()
            job["remote"] = "remote" in mode or "work from home" in mode
            jobs.append(job)
        return jobs
    finally:
        wb.close()


def read_all(sites: list[config.Site] | None = None) -> dict[str, list[dict]]:
    return {s.name: read_site(s) for s in (sites or config.SITES)}
