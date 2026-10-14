import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pipeline import build, config, runner
from pipeline.dedupe import cities, company_tokens, dedupe, group, rank_key, same_company, title_key
from pipeline.load import read_site
from pipeline.route import load_patterns, roles_for, route

TODAY = date(2026, 10, 12)
config.ROLE_SCOPE = "all"          # the default scope is "tech"; most tests exercise every role
PATTERNS = load_patterns()


def job(site, title, company="Acme", location="Bengaluru", url=None, **kw):
    base = {"site": site, "url": url or f"https://{site}.example/{abs(hash((site, title, company, location)))}", "title": title,
            "company": company, "location": location, "more_locations": "", "experience": "", "salary": "", "work_mode": "", "skills": "",
            "job_type": "", "posted_date": date(2026, 10, 10), "first_seen": date(2026, 10, 11), "role_hint": "",
            "bengaluru": "bengaluru" in location.lower() or "bangalore" in location.lower(), "remote": False}
    base.update(kw)
    return base


# ---------------------------------------------------------------- routing
def test_roles_for_titles():
    assert roles_for("Senior Data Scientist", PATTERNS) == ["Data Scientist"]
    assert roles_for("Lead Machine Learning Engineer", PATTERNS) == ["ML Engineer"]
    assert set(roles_for("AI Architect", PATTERNS)) == {"Architect", "AI Engineer"}          # one title can fit two roles
    assert "Forward Deployed Engineer" in roles_for("Forward Deployed Engineer - Full Stack", PATTERNS)
    assert roles_for("Senior Legal Counsel", PATTERNS) == [] and roles_for("SDE-1", PATTERNS) == []
    assert "Ayurveda Consultant" in roles_for("Ayurvedic Doctor (BAMS)", PATTERNS)
    assert any(r.startswith("Clinical") for r in roles_for("Clinical Research Associate", PATTERNS))


def test_route_splits_matched_and_unmatched():
    jobs = [job("hirist", "Data Analyst"), job("hirist", "Sales Executive")]
    matched, unmatched = route(jobs, PATTERNS)
    assert [j["title"] for j in matched] == ["Data Analyst"] and [j["title"] for j in unmatched] == ["Sales Executive"]
    assert matched[0]["roles"] == ["Data Analyst"] and unmatched[0]["roles"] == []


# ---------------------------------------------------------------- de-duplication
def test_company_names_compare_without_legal_suffixes():
    assert company_tokens("Tredence Analytics Solutions Private Limited") == ("tredence", "analytics", "solutions")
    assert same_company(company_tokens("Tredence"), company_tokens("Tredence Analytics Solutions Private Limited"))
    assert same_company(company_tokens("Acme Pvt. Ltd."), company_tokens("ACME (India)"))
    assert not same_company(company_tokens("Acme"), company_tokens("Beta Acme"))                 # a prefix of words only
    assert not same_company((), company_tokens("Acme")) and not same_company((), ())             # no company: never merged by title


def test_title_and_city_normalising():
    assert title_key("Sr. Data Scientist (NLP) ") == title_key("Senior Data Scientist - NLP")
    assert cities("Bengaluru, Karnataka, India") == cities("Bangalore")
    assert cities("India") == frozenset() and cities("Remote") == frozenset()
    assert cities("Gurgaon/Gurugram") == frozenset({"gurugram"})


def test_same_opening_on_two_sites_is_one_row():
    a = job("naukri", "Senior Data Scientist", "Acme Pvt Ltd", "Bengaluru", experience="5-8 Yrs")
    b = job("linkedin", "Senior Data Scientist", "Acme", "Bengaluru, Karnataka, India", salary="")
    c = job("hirist", "Senior Data Scientist", "Acme Private Limited", "Bangalore", salary="20-30 LPA")
    other_city = job("naukri", "Senior Data Scientist", "Acme", "Pune")
    other_title = job("naukri", "Data Engineer", "Acme", "Bengaluru")
    merged = dedupe([a, b, c, other_city, other_title])
    assert len(merged) == 3
    top = next(m for m in merged if len(m["sites"]) == 3)
    assert top["sites"] == ["hirist", "naukri", "linkedin"]                  # richest site first
    assert top["salary"] == "20-30 LPA" and top["experience"] == "5-8 Yrs"   # each field from the first site that has it
    assert top["url"] == c["url"] and set(top["other_links"]) == {a["url"], b["url"]} and len(top["links"]) == 3


def test_unknown_or_generic_place_overlaps_but_other_city_does_not():
    a = job("linkedin", "Data Analyst", "Acme", "India")
    b = job("naukri", "Data Analyst", "Acme", "Pune")
    c = job("hirist", "Data Analyst", "Acme", "Chennai")
    assert len(group([a, b])) == 1 and len(group([b, c])) == 2


def test_no_company_is_never_merged_by_title_but_same_link_is():
    a, b = job("hirist", "Data Analyst", "", "Pune"), job("naukri", "Data Analyst", "", "Pune")
    assert len(group([a, b])) == 2
    same = job("naukri", "Data Analyst", "", "Pune", url=a["url"])
    assert len(group([a, same])) == 1


def test_merge_keeps_dates_and_roles():
    a = job("naukri", "AI Architect", "Acme", "Pune", posted_date=date(2026, 10, 9), first_seen=date(2026, 10, 12))
    b = job("hirist", "AI Architect", "Acme", "Pune", posted_date=date(2026, 10, 8), first_seen=date(2026, 10, 11))
    route([a, b], PATTERNS)
    m = dedupe([a, b])[0]
    assert m["posted_date"] == date(2026, 10, 8) and m["first_seen"] == date(2026, 10, 11) and set(m["roles"]) == {"Architect", "AI Engineer"}


def test_ranking_bengaluru_then_remote_then_rest_newest_first():
    rows = [
        {"bengaluru": False, "remote": False, "posted_date": date(2026, 10, 12), "title": "d"},
        {"bengaluru": False, "remote": True, "posted_date": date(2026, 10, 1), "title": "c"},
        {"bengaluru": True, "remote": False, "posted_date": date(2026, 10, 2), "title": "b"},
        {"bengaluru": True, "remote": False, "posted_date": date(2026, 10, 9), "title": "a"},
    ]
    assert [r["title"] for r in sorted(rows, key=rank_key)] == ["a", "b", "c", "d"]


# ---------------------------------------------------------------- reading a site workbook and building the combined one
def _site_workbook(path: Path, sheet: str, rows: list[dict]) -> config.Site:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Title", "Company", "Role", "Experience", "Location", "Work mode", "Posted date", "First seen", "Link", "Status", "Notes"])
    for r in rows:
        ws.append([r["title"], r["company"], "", r.get("experience", ""), r["location"], r.get("work_mode", ""), r["posted"], r["seen"], "Open", "New", ""])
        ws.cell(row=ws.max_row, column=9).hyperlink = r["url"]
    wb.save(path)
    return config.Site(sheet.lower(), sheet, path, (), 1)


def test_read_site_reads_columns_by_header_and_links_from_hyperlinks(tmp_path):
    site = _site_workbook(tmp_path / "S.xlsx", "Hirist", [
        {"title": "Data Scientist", "company": "Acme", "location": "Bangalore", "work_mode": "Work from home", "posted": date(2026, 10, 10),
         "seen": date(2026, 10, 11), "url": "https://x/1"}])
    jobs = read_site(site)
    assert len(jobs) == 1 and jobs[0]["url"] == "https://x/1" and jobs[0]["bengaluru"] and jobs[0]["remote"]
    assert jobs[0]["posted_date"] == date(2026, 10, 10) and jobs[0]["site"] == "hirist"
    assert read_site(site, tmp_path / "missing.xlsx") == []


def test_build_makes_tabs_keeps_user_edits_and_marks_new_today(tmp_path):
    path = tmp_path / "Job_Radar.xlsx"
    first = {"hirist": [job("hirist", "Data Scientist", "Acme", "Bengaluru", first_seen=date(2026, 10, 11), url="https://h/1"),
                        job("hirist", "Sales Manager", "Beta", "Pune", url="https://h/2")],
             "naukri": [job("naukri", "Data Scientist", "Acme Pvt Ltd", "Bengaluru", first_seen=date(2026, 10, 11), url="https://n/1"),
                        job("naukri", "Business Analyst", "Gamma", "Delhi", first_seen=date(2026, 10, 11), url="https://n/2")]}
    written, stats = build.build(first, path, TODAY)
    assert written == path and stats["unique"] == 3 and stats["duplicates_removed"] == 1 and stats["matched"] == 2 and stats["unmatched"] == 1
    wb = load_workbook(path)
    assert wb.sheetnames[:2] == ["New today", "All"] and "Unmatched" in wb.sheetnames and "Data Scientist" in wb.sheetnames
    ws = wb["All"]
    assert [c.value for c in ws[1]] == build.HEADERS
    assert ws.cell(row=2, column=2).value == "Data Scientist"                                   # Bengaluru ranks first
    assert ws.cell(row=2, column=build.HEADERS.index("Sites") + 1).value == "hirist, naukri"
    assert ws.cell(row=2, column=build.HEADERS.index("Link") + 1).hyperlink.target == "https://h/1"
    assert ws.cell(row=2, column=build.HEADERS.index("Other links") + 1).value == "https://n/1"
    assert wb["New today"].max_row == 1                                                         # first seen was yesterday
    # the owner edits a role tab; a new run brings a new job and the same job again under a different link
    wb["Data Scientist"].cell(row=2, column=build.HEADERS.index("Status") + 1).value = "Applied"
    wb["Data Scientist"].cell(row=2, column=build.HEADERS.index("Notes") + 1).value = "call Monday"
    wb.save(path)
    second = {"hirist": first["hirist"] + [job("hirist", "ML Engineer", "Delta", "Bengaluru", first_seen=TODAY, url="https://h/3")],
              "naukri": [job("naukri", "Data Scientist", "Acme Pvt Ltd", "Bengaluru", first_seen=TODAY, url="https://n/1")]}
    _, stats = build.build(second, path, TODAY)
    wb = load_workbook(path)
    rows = {r[1]: r for r in wb["All"].iter_rows(min_row=2, values_only=True)}
    ds = rows["Data Scientist"]
    assert ds[build.HEADERS.index("Status")] == "Applied" and ds[build.HEADERS.index("Notes")] == "call Monday"
    assert ds[build.HEADERS.index("First seen")].date() == date(2026, 10, 11)                   # not reset
    assert stats["new_today"] == 1 and [r[1] for r in wb["New today"].iter_rows(min_row=2, values_only=True)] == ["ML Engineer"]
    assert rows["ML Engineer"][build.HEADERS.index("Status")] == "New"


def test_build_writes_a_copy_when_excel_has_the_file_open(tmp_path, monkeypatch):
    path = tmp_path / "Job_Radar.xlsx"
    real = Workbook.save
    calls = {"n": 0}

    def locked(self, filename):
        calls["n"] += 1
        if Path(filename) == path:
            raise PermissionError
        return real(self, filename)

    monkeypatch.setattr(Workbook, "save", locked)
    written, _ = build.build({"hirist": [job("hirist", "Data Analyst", url="https://h/1")]}, path, TODAY)
    assert written != path and written.name.startswith("Job_Radar_") and written.exists()


# ---------------------------------------------------------------- running the sites
def test_command_and_site_list():
    site = config.SITE_BY_NAME["hirist"]
    cmd = runner.command(site)
    assert cmd[:4] == [sys.executable, "-m", "hirist", "search"] and "--workers" in cmd
    assert {s.name for s in config.SITES} == {"hirist", "iimjobs", "cutshort", "wellfound", "indeed", "naukri", "linkedin"}


def _fake_site(tmp_path, name, body):
    pkg = tmp_path / name
    pkg.mkdir()
    (pkg / "__init__.py").write_text("")
    (pkg / "__main__.py").write_text(body)
    return config.Site(name, name.title(), tmp_path / "x.xlsx", (), 1)


def test_run_sites_in_parallel_and_report_each(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    ok = _fake_site(tmp_path, "okaysite", "print('3 new, 4 already known. 7 in Okay.xlsx')\n")
    check = _fake_site(tmp_path, "checksite", "print('2 new, 0 already known. 2 in C.xlsx')\nprint('!! Hirist is asking for the human check')\n")
    boom = _fake_site(tmp_path, "boomsite", "import sys\nprint('Traceback')\nsys.exit(1)\n")
    done = []
    results = runner.run_sites([ok, check, boom], timeout_min=2, on_done=done.append, retry=False)
    by = {r["site"]: r for r in results}
    assert by["okaysite"]["status"] == "ok" and (by["okaysite"]["new"], by["okaysite"]["total"]) == (3, 7)
    assert by["checksite"]["status"] == "needs you" and by["checksite"]["problems"]
    assert by["boomsite"]["status"] == "failed" and by["boomsite"]["exit_code"] == 1
    assert len(done) == 3 and all(Path(r["log"]).exists() for r in results)


def test_a_site_running_too_long_is_stopped_and_the_others_finish(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    quick = _fake_site(tmp_path, "quicksite", "print('1 new, 0 already known. 1 in Q.xlsx')\n")
    slow = _fake_site(tmp_path, "slowsite", "import time\ntime.sleep(60)\n")
    results = runner.run_sites([quick, slow], timeout_min=0.03, retry=False)                     # 1.8 s limit; the slow site would sleep for 60 s
    by = {r["site"]: r for r in results}
    assert by["quicksite"]["status"] == "ok" and by["slowsite"]["status"] == "timed out"


def test_scheduler_script_uses_the_pipeline_and_start_when_available(monkeypatch):
    from pipeline import schedule
    seen = {}

    def fake(script):
        seen["script"] = script
        return subprocess.CompletedProcess([], 0, "registered", "")

    monkeypatch.setattr(schedule, "_powershell", fake)
    ok, _ = schedule.register("09:00")
    s = seen["script"]
    assert ok and "-m pipeline run" in s and "-StartWhenAvailable" in s and "-LogonType Interactive" in s and config.TASK_NAME in s
    assert "AddHours(9).AddMinutes(0)" in s and str(config.ROOT) in s



# ---------------------------------------------------------------- lanes, memory, stalls, retry
def _setup(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "ROOT", tmp_path)
    monkeypatch.setattr(config, "LOGS_DIR", tmp_path / "logs")
    monkeypatch.setattr(runner, "POLL_SECONDS", 0.05)
    monkeypatch.setattr(runner, "stop_leftover_chromes", lambda since: 0)
    monkeypatch.setattr(runner, "free_mb", lambda: 10 ** 6)


def test_lanes_follow_the_selected_sites():
    lanes = config.lanes_for(list(config.SITES))
    assert [[s.name for s in lane] for lane in lanes] == config.LANES
    only = config.lanes_for([config.SITE_BY_NAME["iimjobs"], config.SITE_BY_NAME["linkedin"], config.SITE_BY_NAME["hirist"]])
    assert [[s.name for s in lane] for lane in only] == [["linkedin"], ["hirist", "iimjobs"]]


def test_sites_in_a_lane_run_one_after_another_and_lanes_side_by_side(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    first = _fake_site(tmp_path, "firstsite", "import time\ntime.sleep(1.0)\nprint('1 new, 0 already known. 1 in F.xlsx')\n")
    second = _fake_site(tmp_path, "secondsite", "print('1 new, 0 already known. 1 in S.xlsx')\n")
    other = _fake_site(tmp_path, "othersite", "print('1 new, 0 already known. 1 in O.xlsx')\n")
    order = []
    results = runner.run_sites([first, second, other], lanes=[[first, second], [other]], timeout_min=2, on_done=lambda r: order.append(r["site"]))
    assert order == ["othersite", "firstsite", "secondsite"]                 # the other lane did not wait for the first lane
    assert [r["site"] for r in results] == ["firstsite", "secondsite", "othersite"] and all(r["status"] == "ok" for r in results)


def test_a_site_waits_for_memory_and_is_skipped_then_tried_again_at_the_end(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    state = {"free": 100}
    monkeypatch.setattr(runner, "free_mb", lambda: state["free"])
    site = _fake_site(tmp_path, "memsite", "print('1 new, 0 already known. 1 in M.xlsx')\n")

    def on_done(r):
        state["free"] = 10 ** 6                                              # memory comes back after the site was skipped
    results = runner.run_sites([site], lanes=[[site]], timeout_min=2, min_free=3000, memory_wait_min=0.02, on_done=on_done)
    r = results[0]
    assert r["status"] == "ok" and r["retried"] and r["first_status"] == "skipped"


def test_a_site_with_a_silent_log_is_stopped_as_stalled(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    quiet = _fake_site(tmp_path, "quietsite", "import time\ntime.sleep(60)\n")
    results = runner.run_sites([quiet], lanes=[[quiet]], timeout_min=5, stall_sec=1.5, max_reruns=0, retry=False)
    assert results[0]["status"] == "stalled"


def test_a_failed_site_is_retried_once_but_a_site_that_needs_you_is_not(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    flag = tmp_path / "flag.txt"
    flaky = _fake_site(tmp_path, "flakysite", f"import sys, pathlib\nf = pathlib.Path(r'{flag}')\nif f.exists():\n    print('2 new, 0 already known. 2 in F.xlsx')\nelse:\n    f.write_text('x'); sys.exit(1)\n")
    check = _fake_site(tmp_path, "needsite", "print('0 new, 0 already known. 0 in N.xlsx')\nprint('!! asking for the human check')\n")
    results = runner.run_sites([flaky, check], lanes=[[flaky], [check]], timeout_min=2)
    by = {r["site"]: r for r in results}
    assert by["flakysite"]["status"] == "ok" and by["flakysite"]["retried"] and by["flakysite"]["first_status"] == "failed"
    assert by["needsite"]["status"] == "needs you" and not by["needsite"].get("retried")


def test_free_memory_is_read():
    assert runner.free_mb() > 0


def test_a_stuck_site_is_started_again_at_once_and_finishes(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    flag = tmp_path / "stuck_once.txt"
    body = (f"import time, pathlib\nf = pathlib.Path(r'{flag}')\nif not f.exists():\n    f.write_text('x')\n    print('starting', flush=True)\n    time.sleep(60)\n"
            "print('4 new, 0 already known. 4 in S.xlsx')\n")
    stuck = _fake_site(tmp_path, "stucksite", body)
    follower = _fake_site(tmp_path, "followsite", "print('1 new, 0 already known. 1 in F.xlsx')\n")
    seen = []
    results = runner.run_sites([stuck, follower], lanes=[[stuck, follower]], timeout_min=5, stall_sec=1.5,
                               on_done=lambda r: seen.append((r["site"], r["status"])), retry=False)
    by = {r["site"]: r for r in results}
    assert by["stucksite"]["status"] == "ok" and by["stucksite"]["reruns"] == 1 and by["stucksite"]["total"] == 4
    assert by["followsite"]["status"] == "ok"
    assert seen[0][0] == "stucksite" and seen[0][1].startswith("stuck, starting again") and seen[-1] == ("followsite", "ok")


def test_a_site_that_keeps_getting_stuck_gives_up_after_the_reruns(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    quiet = _fake_site(tmp_path, "alwaysstuck", "import time\ntime.sleep(60)\n")
    r = runner.run_sites([quiet], lanes=[[quiet]], timeout_min=5, stall_sec=1.0, max_reruns=2, retry=False)[0]
    assert r["status"] == "stalled" and r["reruns"] == 2


def test_series_lane_and_site_report_with_error_log(tmp_path, monkeypatch):
    from pipeline.report import RunReport
    lanes = config.lanes_for(list(config.SITES), series=True)
    assert len(lanes) == 1 and [s.name for s in lanes[0]] == config.SERIES_ORDER
    _setup(tmp_path, monkeypatch)
    bad = _fake_site(tmp_path, "badsite", "print('window 24 h')\nprint('3 new, 2 already known. 5 in B.xlsx')\nprint('!! Something broke on page 4')\n")
    run_dir = tmp_path / "run"
    report = RunReport(run_dir)
    texts = []
    results = runner.run_sites([bad], lanes=[[bad]], timeout_min=2, retry=False, logs_dir=run_dir,
                               on_done=lambda r: texts.append(report.event(r)))
    assert results[0]["status"] == "finished with errors"
    assert "3 new, 2 already known, 5 jobs" in texts[0] and "Something broke on page 4" in texts[0]
    assert "Something broke" in (run_dir / "errors.log").read_text(encoding="utf-8")
    assert "badsite" in (run_dir / "run_report.md").read_text(encoding="utf-8") and (run_dir / "badsite_report.md").exists()
    assert report.event(dict(results[0], status="stuck, starting again (1 of 2)")) is None


# ---------------------------------------------------------------- tech scope
def test_tech_scope_parks_non_tech_roles(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "ROLE_SCOPE", "tech")
    assert {p.role for p in load_patterns()} == config.role_names(True) and "Ayurveda Consultant" not in config.role_names(True)
    jobs = {"indeed": [job("indeed", "Data Scientist", role_hint="Data Scientist"),
                       job("indeed", "Ayurvedic Doctor (BAMS)", role_hint="Ayurveda Consultant"),
                       job("indeed", "Office Assistant", role_hint="Clinical Research; Drug Inspector"),
                       job("indeed", "Sales Manager", role_hint="Data Analyst")]}
    matched, unmatched, stats = build.combine(jobs)
    assert [j["title"] for j in matched] == ["Data Scientist"]
    assert [j["title"] for j in unmatched] == ["Sales Manager"]          # found by a tech search but fits no role: still shown
    assert stats["parked_non_tech"] == 2


def test_scoped_input_keeps_only_tech_roles(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "ROLE_SCOPE", "tech")
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    path = config.scoped_input()
    rows = path.read_text(encoding="utf-8").splitlines()
    assert rows[0].startswith("Role,") and len(rows) == 9 and not any("Ayurveda" in r or "Clinical" in r or "Drug" in r for r in rows)
    monkeypatch.setattr(config, "ROLE_SCOPE", "all")
    assert config.scoped_input() is None


def test_sites_get_the_scoped_input_file(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "ACTIVE_INPUT", tmp_path / "input_tech.csv")
    cmd = runner.command(config.SITE_BY_NAME["naukri"])
    assert cmd[-2:] == ["--input", str(tmp_path / "input_tech.csv")]


def test_activate_scope_sets_the_input_file_the_sites_get(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(config, "ACTIVE_INPUT", None)
    monkeypatch.setattr(config, "ROLE_SCOPE", "tech")
    line = config.activate_scope()
    assert "tech roles only" in line and config.ACTIVE_INPUT == tmp_path / "input_tech.csv"
    assert runner.command(config.SITE_BY_NAME["linkedin"])[-2:] == ["--input", str(config.ACTIVE_INPUT)]
    monkeypatch.setattr(config, "ROLE_SCOPE", "all")
    assert "all roles" in config.activate_scope() and config.ACTIVE_INPUT is None


def test_a_site_writing_its_excel_file_is_not_stopped_as_stalled(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    body = "import time\nprint('read 5 jobs', flush=True)\nprint('saving to Excel ...', flush=True)\ntime.sleep(4)\nprint('2 new, 3 already known. 5 in S.xlsx')\n"
    site = _fake_site(tmp_path, "savesite", body)
    r = runner.run_sites([site], lanes=[[site]], timeout_min=5, stall_sec=1.0, max_reruns=0, retry=False)[0]
    assert r["status"] == "ok" and r["total"] == 5                  # silent for 4 s with a 1 s stall limit, but it was saving


# ---------------------------------------------------------------- stop / restart one site from the tracker
def _control(cmds):
    """A control source that hands out each batch of commands once, in order."""
    queue = [list(c) for c in cmds]
    return lambda: queue.pop(0) if queue else []


def test_stop_ends_one_running_site_and_the_lane_goes_on(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    slow = _fake_site(tmp_path, "slowone", "import time\nprint('start', flush=True)\ntime.sleep(60)\n")
    nxt = _fake_site(tmp_path, "nextone", "print('3 new, 0 already known. 3 in N.xlsx')\n")
    other = _fake_site(tmp_path, "otherone", "import time\ntime.sleep(2)\nprint('1 new, 0 already known. 1 in O.xlsx')\n")
    notes = []
    results = runner.run_sites([slow, nxt, other], lanes=[[slow, nxt], [other]], timeout_min=5, stall_sec=60, retry=False,
                               control=_control([[], [], [{"site": "slowone", "action": "stop"}]]), on_note=notes.append)
    by = {r["site"]: r for r in results}
    assert by["slowone"]["status"] == "stopped" and "stopped by you" in by["slowone"]["problems"]
    assert by["nextone"]["status"] == "ok" and by["otherone"]["status"] == "ok"        # the next site in the lane and the other lane were not touched
    assert any("slowone: stop requested" in n for n in notes)


def test_stop_skips_a_site_that_is_still_waiting(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    first = _fake_site(tmp_path, "firstone", "import time\ntime.sleep(2)\nprint('1 new, 0 already known. 1 in F.xlsx')\n")
    second = _fake_site(tmp_path, "secondone", "print('9 new, 0 already known. 9 in S.xlsx')\n")
    results = runner.run_sites([first, second], lanes=[[first, second]], timeout_min=5, retry=False,
                               control=_control([[{"site": "secondone", "action": "stop"}]]))
    by = {r["site"]: r for r in results}
    assert by["firstone"]["status"] == "ok" and by["secondone"]["status"] == "stopped" and by["secondone"]["new"] is None


def test_restart_runs_a_running_site_again_with_a_new_log(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    flag = tmp_path / "flag2.txt"
    body = (f"import time, pathlib\nf = pathlib.Path(r'{flag}')\nif not f.exists():\n    f.write_text('x')\n    print('first try', flush=True)\n    time.sleep(60)\n"
            "print('5 new, 1 already known. 6 in R.xlsx')\n")
    site = _fake_site(tmp_path, "restartone", body)
    started = []
    results = runner.run_sites([site], lanes=[[site]], timeout_min=5, stall_sec=60, retry=False, on_start=lambda n, log, k: started.append((n, log.name, k)),
                               control=_control([[], [], [{"site": "restartone", "action": "restart"}]]))
    r = results[0]
    assert r["status"] == "ok" and r["total"] == 6 and r["reruns"] == 1
    assert [s[1:] for s in started] == [("restartone.log", 0), ("restartone_rerun1.log", 1)]


def test_restart_of_a_finished_site_while_other_lanes_run(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    quick = _fake_site(tmp_path, "quickone", "print('2 new, 0 already known. 2 in Q.xlsx')\n")
    longer = _fake_site(tmp_path, "longone", "import time\ntime.sleep(8)\nprint('1 new, 0 already known. 1 in L.xlsx')\n")
    starts = []
    results = runner.run_sites([quick, longer], lanes=[[quick], [longer]], timeout_min=5, retry=False, on_start=lambda n, log, k: starts.append(n),
                               control=_control([[], [], [], [], [{"site": "quickone", "action": "restart"}]]))
    by = {r["site"]: r for r in results}
    assert starts.count("quickone") == 2 and by["quickone"]["status"] == "ok" and by["quickone"]["reruns"] == 1


def test_manual_restart_does_not_use_up_the_automatic_restarts(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    quiet = _fake_site(tmp_path, "quietone", "import time\ntime.sleep(60)\n")
    r = runner.run_sites([quiet], lanes=[[quiet]], timeout_min=5, stall_sec=1.0, max_reruns=1, retry=False,
                         control=_control([[], [{"site": "quietone", "action": "restart"}]]))[0]
    assert r["status"] == "stalled" and r["reruns"] == 2          # 1 by hand + 1 automatic; the one allowed automatic restart was still available after the manual one


def test_control_files_are_read_once_and_deleted(tmp_path):
    (tmp_path / "1_hirist_stop.json").write_text('{"site": "hirist", "action": "stop"}', encoding="utf-8")
    (tmp_path / "2_bad.json").write_text("not json", encoding="utf-8")
    got = runner.read_control(tmp_path)
    assert got == [{"site": "hirist", "action": "stop"}] and runner.read_control(tmp_path) == [] and not list(tmp_path.glob("*.json"))
