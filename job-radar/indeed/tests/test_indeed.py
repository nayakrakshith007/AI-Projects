from datetime import date

import pandas as pd
from openpyxl import load_workbook

from common.inputs import Query
from indeed import config, engine, excel
from indeed.parse import detect_skills, emails_in, experience_from_text, openings_in, phones_in, qualifications, salary_fields, to_jobs

NO_SLEEP = lambda s: None


def job(n, exp=""):
    return {"title": f"Job {n}", "company": "Acme", "location": "KA, India", "experience": exp, "work_mode": "", "salary": "",
            "job_type": "fulltime", "snippet": "text", "posted_text": "", "posted_date": date(2026, 10, 11),
            "url": f"https://in.indeed.com/viewjob?jk={n}"}


def test_experience_from_text():
    assert experience_from_text("Data Scientist + Gurobi (Exp 6-8 yrs)", "") == "6-8 years"
    assert experience_from_text("AI Engineer", "We need 5+ years of experience in ML") == "5+ years"
    assert experience_from_text("AI Engineer", "Minimum 4 years experience with Python") == "4+ years"
    assert experience_from_text("AI Engineer", "Our company has 10 years in business") == ""       # no 'experience' nearby
    assert experience_from_text("", None) == ""


def test_to_jobs_maps_fields_and_keeps_jk_in_the_link():
    df = pd.DataFrame([{
        "job_url": "https://in.indeed.com/viewjob?jk=abc123", "title": "ML Engineer (3-5 yrs)", "company": "Acme",
        "location": "Pune, MH, India", "date_posted": date(2026, 10, 10), "is_remote": True, "min_amount": 600000.0,
        "max_amount": 900000.0, "currency": "INR", "interval": "yearly", "job_type": "fulltime", "description": "Build models.",
    }])
    j = to_jobs(df)[0]
    assert j["url"].endswith("jk=abc123") and j["experience"] == "3-5 years" and j["work_mode"] == "Remote"
    assert j["salary"] == "600,000 - 900,000 INR yearly" and j["posted_date"] == date(2026, 10, 10)


def test_place():
    assert config.place("India") == "India" and config.place("Bengaluru") == "Bengaluru, India"


def test_search_applies_experience_rule_and_dedupes_across_workers():
    fetch = lambda q, hours: ([job(1, "6-8 years"), job(2, "1-2 years"), job(3, "")], False)
    queries = [Query("AI", "a", "India", 5), Query("DS", "b", "India", 5)]
    jobs, err, stats = engine.search(queries, 24, fetch=fetch, sleep=NO_SLEEP, workers=2)
    urls = {j["url"] for j in jobs}
    assert err is None and urls == {"https://in.indeed.com/viewjob?jk=1", "https://in.indeed.com/viewjob?jk=3"}   # 1-2 years dropped
    assert stats == {"fetched": 6, "dropped_experience": 2}
    assert next(j for j in jobs if j["url"].endswith("jk=1"))["role"] == "AI"                                      # first search wins


def test_backs_off_then_succeeds():
    calls, sleeps = [], []

    def fetch(q, hours):
        calls.append(1)
        return ([], True) if len(calls) == 1 else ([job(1)], False)

    jobs, err, _ = engine.search([Query("R", "k", "India", 0)], 24, fetch=fetch, sleep=sleeps.append, workers=1)
    assert len(jobs) == 1 and err is None and len(calls) == 2
    assert any(config.BACKOFF[0] <= s <= config.BACKOFF[1] for s in sleeps)


def test_gives_up_and_stops_the_run():
    queries = [Query("R", f"k{i}", "India", 0) for i in range(8)]
    seen = []

    def fetch(q, hours):
        seen.append(q.keyword)
        return [], True

    jobs, err, _ = engine.search(queries, 24, fetch=fetch, sleep=NO_SLEEP, workers=1)
    assert jobs == [] and "gave up" in err and "stopped early" in err
    assert len(set(seen)) < 8


def test_excel_adds_dedupes_and_keeps_user_fields(tmp_path):
    path = tmp_path / "Indeed.xlsx"
    j = dict(job(1, "6-8 years"), role="AI Engineer", term="ai engineer")
    assert excel.save([j], path, date(2026, 10, 11))[1:] == (1, 0, 1)
    wb = load_workbook(path)
    wb["Indeed"].cell(row=2, column=excel.col("Status")).value = "Applied"
    wb.save(path)
    _, n_new, n_upd, total = excel.save([j, dict(j, url=j["url"] + "9")], path, date(2026, 10, 12))
    assert (n_new, n_upd, total) == (1, 1, 2)
    ws = load_workbook(path)["Indeed"]
    assert "Applied" in [ws.cell(row=r, column=excel.col("Status")).value for r in (2, 3)]


def test_hours_zero_means_no_time_filter(monkeypatch):
    import sys
    import types
    seen = {}

    def fake_scrape(**kw):
        seen.update(kw)
        return pd.DataFrame()

    monkeypatch.setitem(sys.modules, "jobspy", types.SimpleNamespace(scrape_jobs=fake_scrape))
    engine.fetch_query(Query("R", "k", "India", 0), 0)
    assert "hours_old" not in seen and seen["results_wanted"] == config.RESULTS_NO_WINDOW
    engine.fetch_query(Query("R", "k", "India", 0), 24)
    assert seen["hours_old"] == 24 and seen["results_wanted"] == config.RESULTS


def test_more_experience_forms():
    assert experience_from_text("", "Experience: 2-4 years in clinical research") == "2-4 years"
    assert experience_from_text("", "Experience required: 3-5") == "3-5 years"
    assert experience_from_text("", "Freshers can apply") == "0 years (fresher)"
    assert experience_from_text("CRA (2 to 4 years)", "") == "2-4 years"


def test_qualifications():
    assert qualifications("BAMS with MD (Ayurveda) required") == "BAMS, MD/MS Ayurveda"
    assert qualifications("B.Pharm or M.Pharm; Diploma holders welcome") == "B.Pharm, M.Pharm, Diploma"
    assert qualifications("Any graduate") == "Graduate"
    assert qualifications("Strong communication") == ""


def test_detect_skills():
    assert detect_skills("Python, SQL and PyTorch; Docker") == "python, sql, pytorch, docker"
    assert "pharmacovigilance" in detect_skills("Experience in Pharmacovigilance and ICH-GCP")
    assert detect_skills("Cooking") == ""


def test_contacts_and_small_facts():
    text = "Walk-in on Monday. Call 98765 43210 or +91-9123456789, mail hr@clinic.in. 5 vacancies."
    assert emails_in(text, "HR@Clinic.in") == "hr@clinic.in"
    assert phones_in(text) == "9876543210, 9123456789"
    assert openings_in(text) == "5"
    j = to_jobs(pd.DataFrame([{"job_url": "https://in.indeed.com/viewjob?jk=1", "title": "Doctor", "description": text}]))[0]
    assert j["walk_in"] == "Yes" and j["openings"] == "5" and j["phones"] == "9876543210, 9123456789"


def test_salary_fields():
    f = salary_fields({"min_amount": 600000.0, "max_amount": 1200000.0, "currency": "INR", "interval": "yearly"})
    assert f["salary"] == "600,000 - 1,200,000 INR yearly" and f["salary_min_month"] == 50000 and f["salary_max_month"] == 100000
    f = salary_fields({"min_amount": 20000.0, "max_amount": None, "currency": "INR", "interval": "monthly"})
    assert f["salary_min_month"] == 20000 and f["salary_max_month"] is None
    assert salary_fields({})["salary"] == ""


def test_to_jobs_extracts_company_details_and_full_description():
    df = pd.DataFrame([{
        "job_url": "https://in.indeed.com/viewjob?jk=zz", "job_url_direct": "https://corp.example/apply", "title": "Clinical Research Associate",
        "company": "Acme Pharma", "location": "Pune, MH, India", "date_posted": date(2026, 10, 10), "is_remote": False,
        "description": "Need 2-4 years experience. B.Pharm or M.Sc. Monitoring and GCP.", "company_rating": 4.2, "company_reviews_count": 120.0,
        "company_industry": "Pharmaceutical", "company_num_employees": "201 to 500", "company_url_direct": "https://acme.example",
        "company_url": "https://in.indeed.com/cmp/Acme", "company_revenue": "less than $1M (USD)", "company_addresses": "Pune",
    }])
    j = to_jobs(df)[0]
    assert j["experience"] == "2-4 years" and (j["experience_min"], j["experience_max"]) == (2, 4)
    assert j["qualification"] == "B.Pharm, M.Sc" and "gcp" in j["skills"] and "monitoring" in j["skills"]
    assert j["apply_link"] == "https://corp.example/apply" and j["company_website"] == "https://acme.example"
    assert j["company_rating"] == 4.2 and j["company_size"] == "201 to 500" and "Pune" in j["company_info"]
    assert j["description"].startswith("Need 2-4 years")


def test_excel_fills_details_into_old_rows_and_keeps_user_fields(tmp_path):
    path = tmp_path / "Indeed.xlsx"
    old = dict(job(1), role="Ayurveda Consultant", term="bams doctor")           # a first run without the new details
    excel.save([old], path, date(2026, 10, 11))
    wb = load_workbook(path)
    wb["Indeed"].cell(row=2, column=excel.col("Notes")).value = "call them"
    wb.save(path)
    richer = dict(old, qualification="BAMS", skills="panchakarma", emails="hr@x.in", apply_link="https://corp.example/apply",
                  description="full text", experience_min=2, experience_max=4)
    _, n_new, n_upd, total = excel.save([richer], path, date(2026, 10, 12))
    assert (n_new, n_upd, total) == (0, 1, 1)
    ws = load_workbook(path)["Indeed"]
    get = lambda name: ws.cell(row=2, column=excel.col(name)).value
    assert get("Qualification") == "BAMS" and get("Emails") == "hr@x.in" and get("Description") == "full text"
    assert get("Notes") == "call them" and get("Role") == "Ayurveda Consultant" and get("Exp min") == 2
    assert ws.cell(row=2, column=excel.col("Apply link")).hyperlink.target == "https://corp.example/apply"
