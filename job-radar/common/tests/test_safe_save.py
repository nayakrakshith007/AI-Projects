import pytest
from openpyxl import Workbook, load_workbook

from common.safe_save import save_workbook


def _wb(value):
    wb = Workbook()
    wb.active["A1"] = value
    return wb


def test_save_replaces_the_file_and_leaves_no_side_file(tmp_path):
    p = tmp_path / "a.xlsx"
    save_workbook(_wb("one"), p)
    save_workbook(_wb("two"), p)
    assert load_workbook(p).active["A1"].value == "two" and not list(tmp_path.glob("*.part"))


def test_a_failed_save_keeps_the_old_file_intact(tmp_path):
    p = tmp_path / "a.xlsx"
    save_workbook(_wb("old"), p)

    class Boom:
        def save(self, path):
            path.write_bytes(b"half a file")          # the side file is written, then the run dies
            raise RuntimeError("stopped")
    with pytest.raises(RuntimeError):
        save_workbook(Boom(), p)
    assert load_workbook(p).active["A1"].value == "old" and not list(tmp_path.glob("*.part"))
