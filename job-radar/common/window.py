"""Which time window a run should fetch: everything since the latest posted date a site's last complete run saw.

A fixed "last 24 hours" loses a day whenever a run is missed. Instead each site keeps a small state file with the
latest posted date it has fully read. The next run starts at midnight of that date (so the same day is read again to
catch late postings; duplicates are skipped by link) and runs up to now. A run that ended with an error (a human
check, a rate-limit stop) does not move the date, so nothing is skipped after a bad day.
"""
import json
import math
from datetime import date, datetime, time as dtime
from pathlib import Path

MIN_HOURS = 24
MAX_HOURS = 720      # 30 days: the longest window Naukri's own filter offers


def window_hours(latest: date | None, now: datetime | None = None, minimum: int = MIN_HOURS, maximum: int = MAX_HOURS) -> int:
    """Hours from midnight of the latest posted date to now, kept between the minimum and the maximum."""
    if latest is None:
        return minimum
    now = now or datetime.now()
    hours = math.ceil((now - datetime.combine(latest, dtime.min)).total_seconds() / 3600)
    return max(minimum, min(maximum, hours))


def latest_in_workbook(path: Path, sheet: str) -> date | None:
    """Newest 'Posted date' in a site's workbook; used once, before a state file exists."""
    if not path.exists():
        return None
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            return None
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows, None)
        if not header or "Posted date" not in header:
            return None
        i = header.index("Posted date")
        dates = [(v.date() if isinstance(v, datetime) else v) for r in rows if (v := r[i]) and isinstance(v, (date, datetime))]
        return max(dates) if dates else None
    finally:
        wb.close()


def read_latest(state_path: Path, workbook_path: Path | None = None, sheet: str | None = None) -> date | None:
    if state_path.exists():
        try:
            return date.fromisoformat(json.loads(state_path.read_text(encoding="utf-8"))["latest_posted_date"])
        except (ValueError, KeyError, json.JSONDecodeError):
            pass
    return latest_in_workbook(workbook_path, sheet) if workbook_path and sheet else None


def resolve_hours(cli_hours: int | None, state_path: Path, workbook_path: Path, sheet: str, now: datetime | None = None,
                  maximum: int = MAX_HOURS) -> tuple[int, str]:
    """The window for this run and a one-line explanation. An explicit --hours always wins."""
    if cli_hours is not None:
        return cli_hours, f"{cli_hours} h (set on the command line)" if cli_hours else "no time filter (set on the command line)"
    latest = read_latest(state_path, workbook_path, sheet)
    hours = window_hours(latest, now, maximum=maximum)
    if latest is None:
        return hours, f"{hours} h (no earlier run found)"
    capped = " (capped at " + str(maximum) + " h)" if hours == maximum and maximum != MAX_HOURS else ""
    return hours, f"{hours} h (from midnight of the latest posted date already read, {latest}){capped}"


def record_run(state_path: Path, jobs: list[dict], error: str | None, previous: date | None = None) -> date | None:
    """After a run: move the latest posted date forward, but only when the run finished without an error."""
    if error or not jobs:
        return previous
    posted = [j["posted_date"] for j in jobs if j.get("posted_date")]
    if not posted:
        return previous
    latest = max(posted + ([previous] if previous else []))
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps({"latest_posted_date": latest.isoformat(), "updated": datetime.now().isoformat(timespec="seconds")}, indent=2),
                          encoding="utf-8")
    return latest
