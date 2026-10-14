import json
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from common.inputs import Query
from iimjobs import config, excel
from iimjobs import search as S
from iimjobs.parse import job_url, parse_json, split_title, title_matches, total_pages

FIX = Path(__file__).parent / "fixtures"
LIST_JSON = json.loads((FIX / "list.json").read_text(encoding="utf-8"))
TODAY = date(2026, 10, 11)


# ---------------------------------------------------------------- parsing
def test_list_json_fields():
    jobs = parse_json(LIST_JSON)
    assert len(jobs) == 6 and total_pages(LIST_JSON) == 6
    first = jobs[0]
    assert first["id"] == "1738421" and first["url"] == "https://www.iimjobs.com/j/fibe-data-scientist-1738421"
    assert (first["company"], first["title"], first["designation"]) == ("Fibe", "Data Scientist", "Data Scientist")
    assert (first["experience"], first["experience_min"], first["experience_max"]) == ("3-5 years", 3, 5)
    assert first["location"] == "Pune" and first["bengaluru"] == "" and first["salary"] == "" and first["work_mode"] == ""
    assert first["posted_at"] == "2026-10-07T17:20+05:30" and first["posted_date"] == date(2026, 10, 7)       # 1791373811405 ms, Indian time
    assert first["skills"][:2] == ["Data Science", "Analytics"]


def test_salary_remote_and_locations():
    jobs = {j["id"]: j for j in parse_json(LIST_JSON)}
    tredence = jobs["1738078"]
    assert tredence["salary"] == "15-45 LPA" and tredence["bengaluru"] == "Yes"
    assert tredence["company"] == "Tredence Analytics Solutions Private Limited" and tredence["title"] == "Lead Data Scientist"
    assert jobs["1735602"]["work_mode"] == "Work from home"
    straive = jobs["1738122"]
    assert straive["locations"] == ["Bangalore", "Hyderabad", "Gurgaon/Gurugram"] and straive["bengaluru"] == "Yes"


def test_title_without_company_prefix_uses_registered_company():
    head = {j["id"]: j for j in parse_json(LIST_JSON)}["1739063"]
    assert (head["company"], head["title"]) == ("HiringBlaze", "Head - Data Science - Machine Learning & Neural Networks")
    assert split_title("Acme - Data Scientist - NLP", "Acme Corp") == ("Acme Corp", "Data Scientist - NLP")
    assert split_title("Data Scientist", "") == ("", "Data Scientist")


def test_job_links_follow_the_sites_own_pattern():
    # taken from the real page: these are the hrefs IIMJobs itself shows for these jobs
    assert job_url("AkzoNobel - Senior Data Scientist - Advanced Analytics", 1735602) == "https://www.iimjobs.com/j/akzonobel-senior-data-scientist-advanced-analytics-1735602"
    assert job_url("Head - Data Science - Machine Learning & Neural Networks", 1739063) == \
        "https://www.iimjobs.com/j/head-data-science-machine-learning-and-neural-networks-1739063"
    assert job_url("", 5) == "https://www.iimjobs.com/j/5"


def test_nothing_about_recruiters_is_kept():
    for j in parse_json(LIST_JSON):
        assert not any(k in j for k in ("recruiter", "recruiter_name", "created_by", "createdBy", "createdByAlias"))


def test_empty_or_odd_pages():
    assert parse_json({}) == [] and parse_json({"data": [{"title": "x"}]}) == [] and total_pages({"totalPages": 0}) is None


def test_page_count_comes_from_the_job_count():
    # with the date filter on, IIMJobs says totalPages 1 for 89 jobs (real response) and 0 for 12 jobs
    assert total_pages({"totalJobs": 89, "limit": 50, "totalPages": 1}) == 2
    assert total_pages({"totalJobs": 12, "limit": 50, "totalPages": 0}) == 1
    assert total_pages({"totalJobs": 253, "limit": 50, "totalPages": 5}) == 6      # 253 jobs need 6 pages of 50; the site says 5
    assert total_pages({"totalJobs": 0, "totalPages": 0}) is None and total_pages({"totalPages": 3}) == 3


def test_title_matches_is_a_rough_relevance_check():
    assert title_matches({"title": "Senior Data Scientist", "designation": "", "skills": []}, "Data Scientist")
    assert title_matches({"title": "Lead", "designation": "", "skills": ["Machine Learning Engineer"]}, "Machine Learning Engineer")
    assert not title_matches({"title": "Sales Executive", "designation": "", "skills": ["CRM"]}, "Data Scientist")


# ---------------------------------------------------------------- urls and window
def test_posting_days_and_urls():
    assert [config.posting_days(h) for h in (12, 24, 33, 48, 81, 100, 168, 400, 0)] == [1, 1, 3, 3, 7, 7, 7, 7, 7]    # capped at a week
    assert config.tag_url("data-scientist", 7, 1) == "https://www.iimjobs.com/k/data-scientist-jobs?posting=7"
    assert config.tag_url("data-scientist", 7, 3).endswith("posting=7&page=3")
    first, second = config.search_url("ai-engineer", 7, 1), config.search_url("ai-engineer", 7, 2)
    assert "/search/ai-engineer-jobs?" in first and "posting=7" in first and "page=0" in first and "page=2" in second


def test_in_window():
    now = datetime(2026, 10, 12, 9, 0)
    assert S.in_window({"posted_at": "2026-10-11T10:00+05:30", "posted_date": date(2026, 10, 11)}, 33, now)
    assert not S.in_window({"posted_at": "2026-10-10T20:00+05:30", "posted_date": date(2026, 10, 10)}, 33, now)     # 37 h ago
    assert S.in_window({"posted_at": None, "posted_date": None}, 33, now) and S.in_window({"posted_at": "2000-01-01T00:00"}, 0, now)


# ---------------------------------------------------------------- planning
def test_plan_builds_search_and_tag_lists():
    tags = S.load_tags()
    queries = [Query("AI Engineer", "AI Engineer", "India", 5), Query("AI Engineer", "LLM Engineer", "India", 5),
               Query("Forward Deployed Engineer", "Forward Deployed Engineer", "India", 5), Query("DS", "Data Scientist", "India", 5),
               Query("Architect", "Solutions Architect", "India", 5)]
    targets, notes = S.plan(queries, tags)
    pairs = [(t.method, t.slug) for t in targets]
    assert ("search", "ai-engineer") in pairs and ("search", "forward-deployed-engineer") in pairs         # search covers every keyword
    assert [p for p in pairs if p[0] == "tag"].count(("tag", "llm")) == 1                                    # shared tag is read once
    assert ("tag", "artificial-intelligence") in pairs and ("tag", "data-scientist") in pairs and ("tag", "solution-architect") in pairs
    assert not any(t.method == "tag" and t.slug in ("-", "forward-deployed-engineer") for t in targets)
    assert any("Forward Deployed Engineer" in n and "no IIMJobs tag page" in n for n in notes)
    only_search, _ = S.plan(queries, tags, methods=("search",))
    assert {t.method for t in only_search} == {"search"}


# ---------------------------------------------------------------- the shared page queue (fake IIMJobs, no Chrome)
def mk(jid, title="Senior Data Scientist", exp=(5, 8), hours_ago=10, skills=("Python",)):
    posted = datetime.now() - timedelta(hours=hours_ago)
    lo, hi = exp
    return {"id": str(jid), "url": f"https://www.iimjobs.com/j/job-{jid}", "title": title, "company": "Acme", "designation": "",
            "experience": f"{lo}-{hi} years" if hi else "", "experience_min": lo, "experience_max": hi, "locations": ["Bangalore"],
            "location": "Bangalore", "bengaluru": "Yes", "work_mode": "", "salary": "", "skills": list(skills),
            "posted_date": posted.date(), "posted_at": posted.isoformat(timespec="minutes")}


class _FakeSession:
    opened = 0

    def __init__(self, **kw):
        _FakeSession.opened += 1

    def page(self):
        return self

    def close(self):
        pass


def _patch(monkeypatch, lists):
    """lists: {(method, slug): [pages...]}, each page is a list of jobs. Returns the list of (method, slug, page) requested."""
    _FakeSession.opened = 0
    requested, lock = [], threading.Lock()

    def read_page(pg, t, days, page):
        with lock:
            requested.append((t.method, t.slug, page))
        time.sleep(0.01)
        pages = lists.get((t.method, t.slug), [])
        if not pages or page > len(pages):
            return [], (len(pages) or None)
        return [dict(j) for j in pages[page - 1]], len(pages)

    monkeypatch.setattr(S, "read_page", read_page)
    monkeypatch.setattr(S, "Session", _FakeSession)
    monkeypatch.setattr(S, "_pause", lambda rng: None)
    return requested


TAGS = {"data scientist": ["data-scientist"], "forward deployed engineer": ["-"]}


def test_pages_are_shared_and_each_read_once(monkeypatch):
    lists = {("search", "data-scientist"): [[mk(10 * p + i) for i in range(3)] for p in range(1, 6)]}
    requested = _patch(monkeypatch, lists)
    jobs, err, stats = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=3, methods=("search",))
    assert err is None and sorted(p for m, s, p in requested) == [1, 2, 3, 4, 5] and len(jobs) == 15
    assert stats["pages"] == 5 and _FakeSession.opened >= 1 and stats["posting_days"] == 7


def test_both_methods_are_merged_and_compared(monkeypatch):
    shared, only_s, only_t = mk(1), mk(2), mk(3)
    lists = {("search", "data-scientist"): [[shared, only_s]],
             ("tag", "data-scientist"): [[shared, only_t, mk(4, title="Sales Manager", skills=("CRM",))]]}
    _patch(monkeypatch, lists)
    jobs, err, stats = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=2)
    by = {j["id"]: j for j in jobs}
    assert err is None and set(by) == {"1", "2", "3", "4"}
    assert by["1"]["methods"] == ["search", "tag"] and by["2"]["methods"] == ["search"] and by["3"]["methods"] == ["tag"]
    assert by["4"]["title_match"] == ""                                                         # the sales job does not look like the role
    row = stats["compare"][0]
    assert (row["search"], row["tag"], row["both"], row["only_search"], row["only_tag"], row["union"]) == (2, 3, 1, 1, 2, 4)
    assert row["tag_match_title"] == 2 and row["union_match_title"] == 3


def test_window_and_experience_rules(monkeypatch):
    lists = {("search", "data-scientist"): [[mk(1, hours_ago=5), mk(2, hours_ago=400), mk(3, exp=(0, 2)), mk(4, exp=(0, 0)), mk(5, exp=(3, 6))]]}
    _patch(monkeypatch, lists)
    jobs, _, stats = S.search([Query("DS", "Data Scientist", "India", 5)], 168, tags=TAGS, workers=1, methods=("search",))
    assert {j["id"] for j in jobs} == {"1", "4", "5"}          # 2 is too old, 3 asks for 0-2 years; no stated experience (4) is kept; 3-6 reaches 5
    assert stats["outside_window"] == 1 and stats["dropped_experience"] == 1


def test_human_check_stops_every_chrome(monkeypatch):
    from iimjobs.session import HumanCheckPending
    requested = _patch(monkeypatch, {("search", "data-scientist"): [[mk(p * 10 + i) for i in range(3)] for p in range(1, 30)]})
    real = S.read_page

    def read_page(pg, t, days, page):
        if page == 4:
            raise HumanCheckPending("IIMJobs is asking for a human check")
        return real(pg, t, days, page)

    monkeypatch.setattr(S, "read_page", read_page)
    jobs, err, _ = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=3, methods=("search",))
    assert "human check" in err and "check_iimjobs.bat" in err and len(requested) < 25 and jobs


def test_all_chromes_failing_does_not_hang(monkeypatch):
    _patch(monkeypatch, {})

    def boom(pg, t, days, page):
        raise RuntimeError("Chrome crashed")

    monkeypatch.setattr(S, "read_page", boom)
    jobs, err, _ = S.search([Query("DS", "Data Scientist", "India", 0)], 168, tags=TAGS, workers=2, methods=("search",))
    assert jobs == [] and "RuntimeError" in err


# ---------------------------------------------------------------- Excel
def test_excel_adds_dedupes_keeps_user_fields_and_fills_details(tmp_path):
    path = tmp_path / "IIMJobs.xlsx"
    job = dict(mk(7), role="Data Scientist", term="Data Scientist", methods=["tag"], title_match="Yes")
    assert excel.save([job], path, TODAY)[1:] == (1, 0, 1)
    wb = load_workbook(path)
    wb["IIMJobs"].cell(row=2, column=excel.col("Notes")).value = "call them"
    wb.save(path)
    richer = dict(job, methods=["search", "tag"], salary="15-45 LPA")
    _, n_new, n_upd, total = excel.save([richer], path, date(2026, 10, 12))
    assert (n_new, n_upd, total) == (0, 1, 1)
    ws = load_workbook(path)["IIMJobs"]
    get = lambda name: ws.cell(row=2, column=excel.col(name)).value
    assert get("Salary") == "15-45 LPA" and get("Found by") == "search + tag" and get("Notes") == "call them" and get("Role") == "Data Scientist"
    assert ws.cell(row=2, column=excel.col("Link")).hyperlink.target == job["url"]
