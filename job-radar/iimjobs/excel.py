"""Write jobs into the IIMJobs tab of iimjobs/output/IIMJobs.xlsx. New jobs are added by link; Role, Status and Notes are kept."""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config
from common.safe_save import save_workbook

SHEET = "IIMJobs"
HEADERS = ["Title", "Company", "Role", "Experience", "Exp min", "Exp max", "Location", "Bengaluru", "Work mode", "Salary", "Skills",
           "Posted date", "Posted at", "Found by", "Title matches", "Search term", "First seen", "Link", "Status", "Notes"]
LINKCOLS = ("Link",)
WIDTHS = {"Title": 42, "Company": 26, "Role": 22, "Experience": 12, "Exp min": 8, "Exp max": 8, "Location": 24, "Bengaluru": 10,
          "Work mode": 14, "Salary": 12, "Skills": 44, "Posted date": 12, "Posted at": 17, "Found by": 12, "Title matches": 9,
          "Search term": 24, "First seen": 12, "Link": 10, "Status": 14, "Notes": 30}
WRAP = {"Notes"}
DATES = {"Posted date", "First seen"}
NUMBERS = {}
USER_OWNED = {"Role", "Status", "Notes", "First seen", "Link"}
REFRESH = tuple(h for h in HEADERS if h not in USER_OWNED and h != "Search term")
STATUSES = ["New", "Interested", "Applied", "Not relevant"]
THIN = Side(style="thin", color="BFC7D1")


def col(name: str) -> int:
    return HEADERS.index(name) + 1


def _read_existing(path: Path) -> dict[str, dict]:
    """Rows of an existing workbook keyed by link. Columns are matched by header name, so older files still load."""
    if not path.exists():
        return {}
    wb = load_workbook(path)
    if SHEET not in wb.sheetnames:
        return {}
    ws = wb[SHEET]
    names = [c.value for c in ws[1]]
    link_idx = names.index("Link") + 1
    rows = {}
    for row_no, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        link = ws.cell(row=row_no, column=link_idx).hyperlink
        if link and link.target:
            rec = {h: None for h in HEADERS}
            rec.update({k: (v.date() if isinstance(v, datetime) else v) for k, v in zip(names, r) if k in rec})
            for name in LINKCOLS:                    # link cells show "Open"; the address is in the hyperlink
                if name in names:
                    cell = ws.cell(row=row_no, column=names.index(name) + 1)
                    rec[name] = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else ""
            rec["url"] = link.target
            rows[link.target] = rec
    return rows


def _to_record(job: dict, today: date) -> dict:
    return {
        "Title": job["title"], "Company": job["company"], "Role": job.get("role", ""), "Experience": job.get("experience", ""),
        "Exp min": job.get("experience_min"), "Exp max": job.get("experience_max"), "Location": job.get("location", ""),
        "Bengaluru": job.get("bengaluru", ""), "Work mode": job.get("work_mode", ""), "Salary": job.get("salary", ""),
        "Skills": ", ".join(job.get("skills") or []), "Posted date": job["posted_date"], "Posted at": job.get("posted_at") or "",
        "Found by": " + ".join(job.get("methods") or []), "Title matches": job.get("title_match", ""),
        "Search term": job.get("term", ""), "First seen": today, "Link": job["url"], "Status": "New", "Notes": "", "url": job["url"],
    }


def merge(existing: dict[str, dict], jobs: list[dict], today: date) -> tuple[dict[str, dict], int, int]:
    """Add new jobs; for known jobs fill in or refresh the details but keep Role, Status, Notes and First seen."""
    out, n_new, n_upd = dict(existing), 0, 0
    for job in jobs:
        rec = _to_record(job, today)
        old = out.get(job["url"])
        if old is None:
            out[job["url"]] = rec
            n_new += 1
        else:
            for k in REFRESH:
                if rec[k] not in (None, ""):
                    old[k] = rec[k]
            if not old.get("Role"):
                old["Role"], old["Search term"] = rec["Role"], rec["Search term"]
            n_upd += 1
    return out, n_new, n_upd


def _write(rows: dict[str, dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ordered = sorted(rows.values(), key=lambda r: (r["Posted date"] or date.min, r["First seen"] or date.min), reverse=True)
    for rec in ordered:
        ws.append([("Open" if rec.get(h) else None) if h in LINKCOLS else rec.get(h) for h in HEADERS])
        for name in LINKCOLS:
            if rec.get(name):
                ws.cell(row=ws.max_row, column=col(name)).hyperlink = rec[name]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            name = HEADERS[c.column - 1]
            if name in LINKCOLS:
                c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
            else:
                c.font = Font(name="Calibri", size=10)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=name in WRAP)
            if name in DATES:
                c.number_format = "yyyy-mm-dd"
            elif name in NUMBERS:
                c.number_format = NUMBERS[name]
    for name, w in WIDTHS.items():
        ws.column_dimensions[get_column_letter(col(name))].width = w
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    letter = get_column_letter(col("Status"))
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    save_workbook(wb, path)


def save(jobs: list[dict], path: Path = config.EXCEL_PATH, today: date | None = None) -> tuple[Path, int, int, int]:
    """Merge jobs into the workbook. If Excel has it open, a timestamped copy is written instead.
    Returns (file written, new, updated, total)."""
    today = today or date.today()
    existing = _read_existing(path)
    rows, n_new, n_upd = merge(existing, jobs, today)
    try:
        _write(rows, path)
        written = path
    except PermissionError:
        written = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        _write(rows, written)
    return written, n_new, n_upd, len(rows)
