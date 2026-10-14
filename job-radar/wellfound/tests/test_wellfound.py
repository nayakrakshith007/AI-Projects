import threading
import time
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from common.inputs import Query
from wellfound import config, excel
from wellfound import search as S
from wellfound.parse import classify_location, is_generic_page, page_title, parse_page, posted_to_date, total_pages

PAGE = (Path(__file__).parent / "fixtures" / "page.html").read_text(encoding="utf-8")
TODAY = date(2026, 10, 11)


# ---------------------------------------------------------------- parsing
def test_parse_page_company_and_job_fields():
    jobs = parse_page(PAGE, TODAY)
    assert len(jobs) == 3                                              # one single-job company, one with two jobs
    first = jobs[0]
    assert first["title"] == "Data Scientist" and first["company"] == "Recro"
    assert first["url"] == "https://wellfound.com/jobs/4809468-data-scientist"
    assert (first["salary"], first["equity"]) == ("₹14L – ₹17L", "No equity")
    assert (first["work_mode"], first["location"], first["job_type"]) == ("In office", "Bengaluru", "Full-time")
    assert first["experience"] == "2+ years" and first["posted_text"] == "6 days ago" and first["posted_date"] == date(2026, 10, 5)
    assert first["company_size"] == "51-200" and "Growth Stage" in first["company_tags"] and first["actively_hiring"] == "Yes"
    assert first["company_page"] == "https://wellfound.com/company/recrons"


def test_multi_job_company_and_extra_locations():
    jobs = parse_page(PAGE, TODAY)
    assert [j["company"] for j in jobs[1:]] == ["AuxoAI", "AuxoAI"]
    assert jobs[1]["location"] == "Hyderabad" and jobs[1]["more_locations"] == 3 and jobs[1]["experience"] == "5+ years"


def test_onsite_or_remote_mode_is_split_from_the_city():
    first = parse_page(PAGE.replace("In office", "Onsite or remote", 1), TODAY)[0]
    assert (first["work_mode"], first["location"]) == ("Onsite or remote", "Bengaluru")


def test_posted_to_date():
    assert posted_to_date("2 days ago", TODAY) == date(2026, 10, 9)
    assert posted_to_date("3 weeks ago", TODAY) == date(2026, 9, 20)
    assert posted_to_date("1 month ago", TODAY) == date(2026, 9, 11)
    assert posted_to_date("today", TODAY) == TODAY and posted_to_date("", TODAY) is None


def test_generic_page_is_detected():
    assert not is_generic_page(page_title(PAGE), "data-scientist")
    assert is_generic_page("India Jobs - Tech & Startup Jobs | Wellfound", "forward-deployed-engineer")
    assert not is_generic_page("Artificial Intelligence Engineer (AI) Jobs in India - 2027 | Wellfound", "artificial-intelligence-engineer")


def test_total_pages_reads_the_page_n_of_text():
    assert total_pages("<div><span>Page <!-- -->1<!-- --> of <!-- -->24</span></div>") == 24
    assert total_pages("<p>no paging text</p>") is None


def test_location_rule():
    assert classify_location("Bengaluru", "In office", "india") == "India"
    assert classify_location("Bangalore Urban", "", "india") == "India"
    assert classify_location("Onsite or remote • Bengaluru", "", "remote") == "India"
    assert classify_location("India", "Remote only", "remote") == "Remote - India"
    assert classify_location("", "", "remote") == "Remote - location not stated"      # remote with no restriction shown
    assert classify_location("Anywhere", "Remote", "remote") == "Remote - anywhere"
    assert classify_location("United States", "Remote only", "remote").startswith("Outside India")
    assert classify_location("Toronto", "Remote", "remote").startswith("Outside India")
    assert classify_location("Dubai", "In office", "india") == "India"                 # listed on the India list: trusted


# ---------------------------------------------------------------- planning and window
def test_plan_maps_keywords_dedupes_slugs_and_skips_pageless_keywords():
    slugs = S.load_slugs()
    queries = [Query("ML Engineer", "Machine Learning Engineer", "India", 5), Query("ML Engineer", "ML Engineer", "India", 5),
               Query("AI Engineer", "LLM Engineer", "India", 5), Query("Forward Deployed Engineer", "Forward Deployed Engineer", "India", 5),
               Query("Data Scientist", "Data Scientist", "India", 5), Query("X", "Some New Role", "India", 5)]
    targets, notes = S.plan(queries, slugs)
    pairs = [(t.slug, t.kind) for t in targets]
    assert pairs.count(("machine-learning-engineer", "india")) == 1          # two keywords, one page list
    assert ("ai-engineer", "remote") in pairs and ("data-scientist", "india") in pairs
    assert ("some-new-role", "india") in pairs                                # unknown keyword falls back to its own slug
    assert not any(t.slug in ("-", "forward-deployed-engineer") for t in targets)
    assert any("Forward Deployed Engineer" in n for n in notes)


def test_in_window():
    now = datetime(2026, 10, 12, 9, 0)
    assert S.in_window({"posted_date": date(2026, 10, 11)}, 33, now)         # from midnight yesterday
    assert not S.in_window({"posted_date": date(2026, 10, 10)}, 33, now)
    assert S.in_window({"posted_date": None}, 33, now)                        # undated is kept
    assert S.in_window({"posted_date": date(2020, 1, 1)}, 0, now)             # hours = 0: no window at all


# ---------------------------------------------------------------- the shared page queue (fake Wellfound, no Chrome)
class _FakeSession:
    opened: list = []

    def __init__(self, **kw):
        self.profile = kw.get("profile_dir")
        _FakeSession.opened.append(self.profile)

    def page(self):
        return self

    def close(self):
        pass


def _page_html(slug, kind, n, total, posted="3 days ago", loc="Bengaluru", mode="In office", title=None):
    """A synthetic role page built from the real fixture: three jobs whose ids are unique per list and page."""
    base = (sum(map(ord, slug)) * 7 + (1 if kind == "india" else 2)) * 1000 + n * 10
    html = PAGE.replace("4809468", str(base + 1)).replace("4670319", str(base + 2)).replace("4092884", str(base + 3))
    for label in ("6 days ago", "1 month ago", "2 months ago"):
        html = html.replace(label, posted)
    html = html.replace("In office", mode).replace("Bengaluru", loc).replace("Hyderabad", loc).replace("Gurgaon", loc)
    html = html.replace("<title>Data Scientist Jobs in India - 2027 | Wellfound</title>",
                        f"<title>{title or slug.replace('-', ' ').title()} Jobs in India - 2027 | Wellfound</title>")
    return html.replace("</body>", f"<span>Page {n} of {total}</span></body>") if total else html


def _patch(monkeypatch, pages_per_list, say_total=True, **page_kw):
    """Fake Wellfound: every list has the given number of pages; returns the list of pages that were requested."""
    _FakeSession.opened = []
    requested, lock = [], threading.Lock()

    def load(pg, url):
        kind = "india" if "/role/l/" in url else "remote"
        slug = url.split("/role/")[1].split("/")[1].split("?")[0]
        n = int(url.split("page=")[1]) if "page=" in url else 1
        total = pages_per_list.get((slug, kind), pages_per_list.get("default", 3))
        with lock:
            requested.append((slug, kind, n))
        time.sleep(0.02)                                          # lets the other Chromes pick up pages too
        return None if n > total else _page_html(slug, kind, n, total if say_total else None, **page_kw)

    monkeypatch.setattr(S, "_load", load)
    monkeypatch.setattr(S, "Session", _FakeSession)
    monkeypatch.setattr(S, "_pause", lambda rng: None)
    return requested


DS = {"data scientist": "data-scientist"}


def test_pages_of_one_list_are_spread_over_several_chromes_and_each_read_once(monkeypatch):
    requested = _patch(monkeypatch, {("data-scientist", "india"): 12, ("data-scientist", "remote"): 1})
    jobs, err, stats = S.search([Query("Data Scientist", "Data Scientist", "India", 0)], 0, slugs=DS, workers=4)
    india = sorted(n for slug, kind, n in requested if kind == "india")
    assert err is None and india == list(range(1, 13))                        # pages 1..12, none twice, none missed
    assert len(set(_FakeSession.opened)) == 4                                  # four Chromes, four different profile folders
    assert sum(1 for c in stats["pages_per_chrome"].values() if c > 0) >= 2    # the list really was shared
    assert stats["pages"] == 13 and len(jobs) == 13 * 3 - 3 * 0 - (0)          # 13 pages x 3 jobs, all unique


def test_a_list_that_does_not_say_its_page_count_is_read_until_it_runs_out(monkeypatch):
    requested = _patch(monkeypatch, {("data-analyst", "india"): 4, ("data-analyst", "remote"): 1}, say_total=False)
    S.search([Query("DA", "Data Analyst", "India", 0)], 0, slugs={"data analyst": "data-analyst"}, workers=2)
    assert sorted(n for s_, k, n in requested if k == "india") == [1, 2, 3, 4, 5]    # page 5 is empty and ends the list


def test_generic_page_list_is_skipped_and_reported(monkeypatch):
    requested = _patch(monkeypatch, {"default": 3})
    monkeypatch.setattr(S, "page_title", lambda html: "India Jobs - Tech & Startup Jobs | Wellfound")
    jobs, err, stats = S.search([Query("FDE", "Forward Deployed Engineer", "India", 0)], 0, slugs={}, workers=2)
    assert jobs == [] and any("generic page" in n for n in stats["notes"])
    assert max(n for s_, k, n in requested) == 1                                       # only page 1 of each list was opened


def test_window_rule_is_applied_to_what_was_read(monkeypatch):
    _patch(monkeypatch, {"default": 1}, posted="6 days ago")
    jobs, _, stats = S.search([Query("DS", "Data Scientist", "India", 5)], 33, slugs=DS, workers=1)
    assert jobs == [] and stats["outside_window"] > 0                                  # "6 days ago" is outside a 33 h window
    jobs, _, _ = S.search([Query("DS", "Data Scientist", "India", 5)], 168, slugs=DS, workers=1)
    assert jobs                                                                        # and inside a 7 day window
    _, _, stats = S.search([Query("DS", "Data Scientist", "India", 5)], 0, slugs=DS, workers=1)
    assert stats["outside_window"] == 0                                                # hours 0: no window


def test_jobs_are_stamped_with_role_eligibility_and_bengaluru(monkeypatch):
    _patch(monkeypatch, {"default": 1}, posted="2 days ago")
    jobs, _, _ = S.search([Query("Data Scientist", "Data Scientist", "India", 0)], 0, slugs=DS, workers=1)
    india = [j for j in jobs if j["source_list"] == "india"]
    assert india and all(j["role"] == "Data Scientist" and j["eligibility"] == "India" for j in india)
    assert all(j["bengaluru"] == "Yes" for j in india if j["location"] == "Bengaluru")


def test_remote_jobs_outside_india_are_dropped_unless_asked(monkeypatch):
    _patch(monkeypatch, {"default": 1}, loc="Toronto", mode="Remote")
    q = [Query("DS", "Data Scientist", "India", 0)]
    kept, _, stats = S.search(q, 0, slugs=DS, workers=1)
    assert all(j["source_list"] == "india" for j in kept) and stats["dropped_location"] > 0      # the remote list's Toronto jobs are dropped
    everything, _, _ = S.search(q, 0, slugs=DS, workers=1, keep_all_locations=True)
    assert any(j["eligibility"].startswith("Outside India") for j in everything)


def test_human_check_stops_all_chromes_and_keeps_what_was_read(monkeypatch):
    from wellfound.session import HumanCheckPending
    requested = _patch(monkeypatch, {"default": 50})
    real = S._load

    def load(pg, url):
        if "page=4" in url:
            raise HumanCheckPending("Wellfound is asking for a human check")
        return real(pg, url)

    monkeypatch.setattr(S, "_load", load)
    jobs, err, _ = S.search([Query("DS", "Data Scientist", "India", 0)], 0, slugs=DS, workers=3)
    assert "human check" in err and "check_wellfound.bat" in err
    assert len(requested) < 60 and len(jobs) > 0                                      # it stopped long before reading 100 pages


def test_all_chromes_failing_does_not_hang(monkeypatch):
    _patch(monkeypatch, {"default": 5})

    def boom(pg, url):
        raise RuntimeError("Chrome crashed")

    monkeypatch.setattr(S, "_load", boom)
    jobs, err, _ = S.search([Query("DS", "Data Scientist", "India", 0)], 0, slugs=DS, workers=2)
    assert jobs == [] and "RuntimeError" in err


# ---------------------------------------------------------------- Excel
def test_excel_adds_dedupes_and_keeps_user_fields(tmp_path):
    path = tmp_path / "Wellfound.xlsx"
    job = dict(parse_page(PAGE, TODAY)[0], role="Data Scientist", term="Data Scientist", eligibility="India", bengaluru="Yes")
    assert excel.save([job], path, TODAY)[1:] == (1, 0, 1)
    wb = load_workbook(path)
    wb["Wellfound"].cell(row=2, column=excel.col("Status")).value = "Applied"
    wb.save(path)
    other = dict(job, url=job["url"] + "9")
    _, n_new, n_upd, total = excel.save([job, other], path, date(2026, 10, 12))
    assert (n_new, n_upd, total) == (1, 1, 2)
    ws = load_workbook(path)["Wellfound"]
    assert "Applied" in [ws.cell(row=r, column=excel.col("Status")).value for r in (2, 3)]
    assert ws.cell(row=2, column=excel.col("Bengaluru")).value == "Yes"
    assert ws.cell(row=2, column=excel.col("Company page")).hyperlink is not None
