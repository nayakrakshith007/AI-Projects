from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from linkedin import config, excel
from linkedin.parse import parse_card, parse_page, posted_to_date

CARD = (Path(__file__).parent / "fixtures" / "card.html").read_text(encoding="utf-8")
TODAY = date(2026, 10, 11)


def test_parse_card_fields():
    job = parse_card(CARD, TODAY)
    assert job["title"] == "Forward Deployed Engineer — Enterprise AI"
    assert job["company"] == "Riyalabs.ai"
    assert job["location"] == "India"
    assert job["posted_text"] == "7 hours ago"
    assert job["posted_date"] == date(2026, 10, 11)
    assert job["url"] == "https://in.linkedin.com/jobs/view/forward-deployed-engineer-%E2%80%94-enterprise-ai-at-riyalabs-ai-4477243510"


def test_parse_page_finds_all_cards():
    assert len(parse_page(CARD + CARD.replace("4477243510", "1"), TODAY)) == 2


def test_posted_to_date():
    assert posted_to_date("7 hours ago", TODAY) == TODAY
    assert posted_to_date("2 days ago", TODAY) == date(2026, 10, 9)
    assert posted_to_date("", TODAY) is None


def test_search_url():
    url = config.search_url("AI Engineer", 24)
    assert "keywords=AI%20Engineer" in url and "location=India" in url and "f_TPR=r86400" in url and "sortBy=DD" in url


def test_excel_adds_dedupes_keeps_user_fields_and_survives_second_save(tmp_path):
    path = tmp_path / "LinkedIn.xlsx"
    job = parse_card(CARD, TODAY)
    job["term"] = "forward deployed engineer"
    assert excel.save([job], path, TODAY)[1:] == (1, 0, 1)

    wb = load_workbook(path)
    wb["LinkedIn"].cell(row=2, column=excel.HEADERS.index("Status") + 1).value = "Applied"
    wb.save(path)

    other = dict(job, url=job["url"] + "-2")
    _, n_new, n_upd, total = excel.save([job, other], path, date(2026, 10, 12))
    assert (n_new, n_upd, total) == (1, 1, 2)
    ws = load_workbook(path)["LinkedIn"]
    statuses = [ws.cell(row=r, column=excel.HEADERS.index("Status") + 1).value for r in (2, 3)]
    assert "Applied" in statuses
